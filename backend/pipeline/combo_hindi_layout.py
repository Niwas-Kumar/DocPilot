"""Layout of the published combined Hindi edition, isolated from Daily CA.

The Word body is the layout authority. Bookmarks locate section endings in
the actual render; promos are added to those paragraphs without reflowing it.
"""
import os
import io
from copy import deepcopy

import fitz
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw

from . import daily_ca_hindi as H, dea_common as C, dea_hindi as D, cover_kit as K
from .convert import _word_lock
from .ooxml import add_anchored_image, get_or_add

PREFIX = "combo_end_"
LEADING = 16.2
FALLBACK_LEADING = 15.2


def _text(el):
    return C.clean(''.join(t.text or '' for t in el.iter(qn('w:t'))))


def _portrait_galleries(doc):
    """Pair secondary portrait infographics in an article in a clear band.

    The introductory illustration remains a float. Only artwork moves;
    paragraph order and all article text are preserved.
    """
    from docx.shared import Pt
    groups, current = [], []
    for b in list(doc.element.body):
        if b.tag == qn('w:tbl') and any(s.get(qn('w:fill')) == H.PEACH for s in b.iter(qn('w:shd'))):
            if current:
                groups.append(current)
            current = []
        if b.tag != qn('w:p'):
            continue
        for a in b.iter():
            if a.tag not in (qn('wp:anchor'), qn('wp:inline')):
                continue
            ex = a.find(qn('wp:extent'))
            if ex is not None:
                w, h = int(ex.get('cx')), int(ex.get('cy'))
                if 0.6 < w / max(h, 1) < 1.05 and w > 180 * 12700:
                    current.append((b, a))
    if current:
        groups.append(current)
    for pictures in groups:
        for offset in range(1, len(pictures) - 1, 2):
            pair = pictures[offset:offset + 2]
            table = doc.add_table(rows=1, cols=2)
            table.autofit = False
            for col in table.columns:
                col.width = Pt(H.COL_W_PT / 2)
            borders = get_or_add(table._tbl.tblPr, 'w:tblBorders')
            for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
                get_or_add(borders, 'w:' + side).set(qn('w:val'), 'nil')
            target = pair[-1][0]
            target.addprevious(table._tbl)
            for cell, (_carrier, anchor) in zip(table.rows[0].cells, pair):
                blip = next(anchor.iter(qn('a:blip')))
                part = doc.part.related_parts[blip.get(qn('r:embed'))]
                with Image.open(io.BytesIO(part.blob)) as im:
                    width = min(255, 273 * im.width / im.height)
                p = cell.paragraphs[0]
                p.alignment = 1
                p.paragraph_format.space_after = Pt(0)
                p.add_run().add_picture(io.BytesIO(part.blob), width=Pt(width))
                drawing = anchor.getparent()
                drawing.getparent().remove(drawing)


def _content_tables(body):
    """Discard input row minima and cell overrides that inflate branded grids."""
    for t in body.iter(qn('w:tbl')):
        rows = t.findall(qn('w:tr'))
        if not rows:
            continue
        if len(rows) <= 2 and any(s.get(qn('w:fill')) == H.RED for s in t.iter(qn('w:shd'))):
            continue  # measured topic and news bars
        grid = t.find(qn('w:tblGrid'))
        if grid is not None and len(grid) == 3:
            widths = [round(H.COL_W_PT * v * 20) for v in (.24, .49, .27)]
            for col, width in zip(grid, widths):
                col.set(qn('w:w'), str(width))
            for row in rows:
                for cell, width in zip(row.findall(qn('w:tc')), widths):
                    get_or_add(get_or_add(cell, 'w:tcPr'), 'w:tcW').set(qn('w:w'), str(width))
        for row in rows:
            pr = row.find(qn('w:trPr'))
            if pr is not None:
                for height in pr.findall(qn('w:trHeight')):
                    pr.remove(height)
                header = pr.find(qn('w:tblHeader'))
                if header is not None and row is rows[0] and grid is not None and len(grid) >= 3:
                    header.set(qn('w:val'), '1')
        for cell in t.iter(qn('w:tc')):
            pr = get_or_add(cell, 'w:tcPr')
            for border in pr.findall(qn('w:tcBorders')):
                pr.remove(border)
            mar = get_or_add(pr, 'w:tcMar')
            for side, value in (('top', 0), ('bottom', 0), ('left', 80), ('right', 80)):
                el = get_or_add(mar, 'w:' + side)
                el.set(qn('w:w'), str(value))
                el.set(qn('w:type'), 'dxa')


def _last_paragraph(block):
    ps = list(block.iter(qn('w:p')))
    return next((p for p in reversed(ps) if C.clean(C.para_text(p))), ps[-1] if ps else None)


def _mark(p, name, ident):
    if p is None:
        return
    p.insert(1 if p.find(qn('w:pPr')) is not None else 0, parse_xml(
        f'<w:bookmarkStart {C.W_NS} w:id="{ident}" w:name="{PREFIX}{name}"/>'))
    p.append(parse_xml(f'<w:bookmarkEnd {C.W_NS} w:id="{ident}"/>'))


def _furniture(doc, workdir, first_number):
    # An image behind the text gives Word the same double frame as the PDF.
    frame = os.path.join(workdir, 'combo_hi_frame.png')
    if not os.path.exists(frame):
        im = Image.new('RGBA', (2448, 3168))
        draw = ImageDraw.Draw(im)
        for r, width in ((H.FRAME_OUTER, 8), (H.FRAME_INNER, 3)):
            draw.rectangle(tuple(round(v * 4) for v in r), outline='#939598', width=width)
        im.save(frame)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    add_anchored_image(hp, sec.header.part, frame, 0, 0, 612, 792,
                       behind=True, name='combo_frame')
    # A real PAGE field in an ellipse, rather than a PDF-only rounded square.
    cv = K.Cover()
    field = K.run('2', 'Arial', 11, color='FFFFFF', bold=True)
    field = field.replace('<w:t xml:space="preserve">2</w:t>',
                          '<w:fldChar w:fldCharType="begin"/></w:r><w:r>'
                          '<w:instrText xml:space="preserve"> PAGE </w:instrText>'
                          '</w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                          '<w:r><w:rPr><w:color w:val="FFFFFF"/><w:b/><w:sz w:val="22"/>'
                          '</w:rPr><w:t>2</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/>')
    cv.textbox((293.2, 749.9, 25.6, 25.6),
               [K.para([field], jc='center', line=13)], anchor='ctr', fill=H.INK,
               insets=(0, 0, 0, 0))
    shape = deepcopy(cv.host[-1])
    for geom in shape.iter(qn('a:prstGeom')):
        geom.set('prst', 'ellipse')
    sec.footer.paragraphs[0]._p.append(shape)
    get_or_add(sec._sectPr, 'w:pgNumType').set(qn('w:start'), str(first_number))


def normalize(path, workdir, fonts, *, editorial=False, topics=(), day_starts=(), first_number=2):
    """Continuous articles, template pitch, shaded shorts and safe images."""
    doc = Document(path)
    body = doc.element.body
    if not editorial:
        _portrait_galleries(doc)
    _content_tables(body)
    blocks = [b for b in body if b.tag in (qn('w:p'), qn('w:tbl'))]
    for p in body.iter(qn('w:p')):
        text = C.clean(C.para_text(p))
        if not text:
            continue
        ppr = get_or_add(p, 'w:pPr')
        sp = get_or_add(ppr, 'w:spacing')
        # Banner titles retain their measured geometry. Ordinary paragraphs
        # use explicit leading so fallback font metrics cannot add 3pt/line.
        tables = [a for a in p.iterancestors() if a.tag == qn('w:tbl')]
        title = any(len(t.findall(qn('w:tr'))) <= 2
                    and len(t.findall(qn('w:tr'))[-1].findall(qn('w:tc'))) <= 2
                    and any(s.get(qn('w:fill')) == H.RED for s in t.iter(qn('w:shd')))
                    for t in tables if t.findall(qn('w:tr')))
        if not title:
            fallback = fonts[1].replace(' ', '').lower() == 'baloo2'
            sp.set(qn('w:line'), str(round((FALLBACK_LEADING if fallback else LEADING) * 20)))
            sp.set(qn('w:lineRule'), 'exact')
            # Adobe's nominal size has a smaller visual footprint than Baloo.
            # Tune only the bundled fallback; keep licensed Adobe at 11.5pt.
            if fallback:
                for r in p.iter(qn('w:r')):
                    rf = r.find(qn('w:rPr') + '/' + qn('w:rFonts'))
                    if rf is not None and rf.get(qn('w:cs'), '').replace(' ', '').lower() == 'baloo2':
                        for tag in ('w:sz', 'w:szCs'):
                            sz = r.find(qn('w:rPr') + '/' + qn(tag))
                            if sz is not None and sz.get(qn('w:val')) in ('23', '22'):
                                sz.set(qn('w:val'), '18' if sz.get(qn('w:val')) == '23' else '20')
                        # CA copy has more inline bold and long factual lists.
                        # A modest width adjustment preserves the reference's
                        # line lengths with the wider bundled face.
                        if not editorial:
                            get_or_add(r.find(qn('w:rPr')), 'w:w').set(qn('w:val'), '92')
        in_cell = any(a.tag == qn('w:tc') for a in p.iterancestors())
        if in_cell and not title:
            for tag in ('w:keepNext', 'w:keepLines'):
                for el in ppr.findall(qn(tag)):
                    ppr.remove(el)
        elif not in_cell:
            C.set_flag(p, 'w:widowControl', True)
        C.order_ppr(ppr)
    # Wide infographics belong in the flow; portrait pictures may wrap text.
    for anchor in list(body.iter(qn('wp:anchor'))):
        ext = anchor.find(qn('wp:extent'))
        if ext is None:
            continue
        w, h = int(ext.get('cx')), int(ext.get('cy'))
        anchor.set('allowOverlap', '0')
        if w / max(h, 1) < 1.3:
            continue
        width = min(H.COL_W_PT * .66, 360)
        height = width * h / w
        inline = parse_xml('<wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" distT="0" distB="0" distL="0" distR="0"/>')
        for tag in ('wp:extent', 'wp:effectExtent', 'wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
            el = anchor.find(qn(tag))
            if el is not None:
                inline.append(deepcopy(el))
        inline.find(qn('wp:extent')).set('cx', str(round(width * 12700)))
        inline.find(qn('wp:extent')).set('cy', str(round(height * 12700)))
        for ex in inline.iter(qn('a:ext')):
            ex.set('cx', str(round(width * 12700)))
            ex.set('cy', str(round(height * 12700)))
        run = anchor.getparent()
        p = next(a for a in run.iterancestors() if a.tag == qn('w:p'))
        # A separate carrier prevents a full-width drawing sharing a text line.
        carrier = parse_xml(f'<w:p {C.W_NS}><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:drawing/></w:r></w:p>')
        H._reset_ppr(carrier, jc='center', before=8, after=8, line=12, rule='auto')
        carrier.find('.//' + qn('w:drawing')).append(inline)
        p.addprevious(carrier)
        run.remove(anchor)
        if not len(run):
            run.getparent().remove(run)
    if editorial:
        for k, start in enumerate(day_starts[1:]):
            if start >= len(topics):
                continue
            boundary = next((i for i, b in enumerate(blocks)
                             if _text(b) == C.clean(topics[start])), None)
            if boundary is not None and boundary > 0:
                previous = next((b for b in reversed(blocks[:boundary]) if _text(b)), None)
                if previous is not None:
                    _mark(_last_paragraph(previous), f'day{k}', 28000 + k)
    else:
        news = False
        for i, b in enumerate(blocks):
            if b.tag != qn('w:tbl'):
                continue
            # A one-cell red bar is the News in Shorts section, while article
            # banners use an icon and a title in two cells.
            cells = list(b.iter(qn('w:tc')))
            rows = b.findall(qn('w:tr'))
            if rows and len(rows[-1].findall(qn('w:tc'))) == 1 and any(s.get(qn('w:fill')) == H.RED for s in b.iter(qn('w:shd'))):
                news = True
                C.break_before(b)
                if i:
                    _mark(_last_paragraph(blocks[i - 1]), 'articles', 27000)
                continue
            if news:
                for tc in cells:
                    H._shading(get_or_add(tc, 'w:tcPr'), H.PEACH_SOFT)
        if blocks:
            last = next((b for b in reversed(blocks) if _text(b)), blocks[-1])
            _mark(_last_paragraph(last), 'ca', 27001)
    _furniture(doc, workdir, first_number)
    doc.save(path)


def marker_pages(path, cancelled=lambda: False):
    """Read actual physical page numbers, with the same Word engine as export."""
    import pythoncom
    import win32com.client
    found = {}
    with _word_lock:
        pythoncom.CoInitialize()
        word = doc = None
        try:
            word = win32com.client.DispatchEx('Word.Application')
            word.Visible, word.DisplayAlerts = False, 0
            doc = word.Documents.Open(os.path.abspath(path), ReadOnly=True, AddToRecentFiles=False)
            doc.Repaginate()
            for bm in doc.Bookmarks:
                if cancelled():
                    raise InterruptedError('stopped')
                if bm.Name.startswith(PREFIX):
                    rng = bm.Range.Duplicate
                    rng.Collapse(0)  # the paragraph may span more than one page
                    found[bm.Name[len(PREFIX):]] = int(rng.Information(3)) - 1
        finally:
            if doc is not None:
                doc.Close(0)
            if word is not None:
                word.Quit()
            pythoncom.CoUninitialize()
    return found


def fill_promos(pdf, ca_pages, ca_markers, dea_markers):
    """The reference places Foundation before Shorts, Srijan at day endings."""
    schedule = []
    if 'articles' in ca_markers and 0 <= ca_markers['articles'] < ca_pages:
        schedule.append((ca_markers['articles'], D.FOUNDATION, D.FOUNDATION_W, D.FOUNDATION_CX, 'articles'))
    if 'ca' in ca_markers and 0 <= ca_markers['ca'] < ca_pages:
        schedule.append((ca_markers['ca'], D.SRIJAN, 356.7, D.SRIJAN_CX, 'ca'))
    schedule += [(ca_pages + p, D.SRIJAN, D.SRIJAN_W, D.SRIJAN_CX, key)
                 for key, p in dea_markers.items() if key.startswith('day')
                 and 0 <= p < len(pdf) - ca_pages - 1]
    placed, used = {}, set()
    for pno, image, width, cx, key in schedule:
        if pno in used or not 0 <= pno < len(pdf):
            continue
        used.add(pno)
        rect = C.place_banner(pdf[pno], image, C.content_bottom(pdf[pno], D.PAGE_FLOOR),
                              D.PAGE_FLOOR, width, cx, D.BANNER_MIN_SCALE, D.BANNER_GAP,
                              min_top=D.BANNER_TOP)
        if rect is not None:
            placed[key] = (image, tuple(rect), pno)
    return placed


def add_promos(path, placed):
    doc = Document(path)
    for bm in list(doc.element.body.iter(qn('w:bookmarkStart'))):
        key = bm.get(qn('w:name'), '').removeprefix(PREFIX)
        if key not in placed:
            continue
        image, rect, _ = placed[key]
        p = bm.getparent()
        tables = [a for a in p.iterancestors() if a.tag == qn('w:tbl')]
        if tables:
            # Word clips page-positioned pictures whose anchor lives in a
            # cell. Use the hairline paragraph after the outermost table.
            table = tables[-1]
            tail = table.getnext()
            if tail is None or tail.tag != qn('w:p') or _text(tail) or C.has_drawing(tail):
                tail = parse_xml(f'<w:p {C.W_NS}><w:pPr/></w:p>')
                table.addnext(tail)
            H._reset_ppr(tail, line=1, rule='exact', before=0, after=0)
            p = tail
        run = add_anchored_image(Paragraph(p, doc), doc.part, image, rect[0], rect[1],
                                 rect[2] - rect[0], rect[3] - rect[1], name='combo_promo_' + key)
        next(run._r.iter(qn('wp:anchor'))).set('layoutInCell', '0')
    doc.save(path)


def add_badge(path, image, rect):
    doc = Document(path)
    p = parse_xml(f'<w:p {C.W_NS}><w:pPr><w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/></w:pPr></w:p>')
    doc.element.body.insert(0, p)
    add_anchored_image(Paragraph(p, doc), doc.part, image,
                       rect.x0, rect.y0, rect.width, rect.height, name='combo_badge')
    doc.save(path)


def add_closing(path):
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Pt
    doc = Document(path)
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    for hf in (sec.header, sec.footer):
        for p in hf.paragraphs:
            p.clear()
    for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(sec, side, Pt(0))
    p = doc.add_paragraph()
    H._reset_ppr(p._p, line=1, rule='exact', after=0)
    add_anchored_image(p, doc.part, D.CLOSING_PAGE, 0, 0, 612, 792, name='combo_closing')
    doc.save(path)
