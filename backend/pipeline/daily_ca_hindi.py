# -*- coding: utf-8 -*-
"""Daily Current Affairs Summary - HINDI edition.

Restyles a raw typed Hindi .docx into the official PW OnlyIAS Hindi design and
exports the branded PDF.  This is a FORMATTING tool: nothing is ever invented,
added or reworded - every glyph in the output came from the input file.

The Hindi template is a different design from the English one (see
``daily_ca.py`` / ``styler.py`` / ``template.py``, which stay untouched):

  * Letter page, margins 34 / 32 / 69 / 74 pt; no Word page border - the faint
    grey double rule the official sample carries is stamped into the PDF
    afterwards (see DRAW_PAGE_FRAME)
  * header  : faint centred watermark (behind text) + PW OnlyIAS logo
  * footer  : three brand buttons (WhatsApp / PW Web-App / Telegram)
  * page no.: white Arial Narrow Bold in a black rounded box, starting at 2
  * closing : a separate numbered page with the published course/book ads
  * topics  : red icon block + peach banner bar (a 1x2 borderless table)
  * bullets : literal Wingdings glyphs with a hanging indent (red circle /
              black square / black tick), never Word's list engine

Everything measured from the official sample
``Output 17 Aug 2026_CA Analysis_UPSC_Hindi.pdf``.
"""
import os
import re
import shutil
import threading

import fitz
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt

from .convert import docx_to_pdf
from .fitter import fit_document, HINDI_SPEC
from .ooxml import (add_anchored_image, drop_google_tab_frontmatter,
                    pt_to_emu, pt_to_twips, set_run_fonts)

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ASSETS = os.path.join(ROOT, "assets", "hindi_ca")

NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"')

# ---------------------------------------------------------------- palette
RED = "CB202F"          # brand red - banners, bullets, table rules
PEACH = "F3CFC3"        # topic banner bar
PEACH_SOFT = "F9E5DD"   # info-box / table body wash
INK = "231F20"          # body + heading text
WHITE = "FFFFFF"
FRAME_GREY = "939598"   # the faint double rule around every page

# ---------------------------------------------------------------- geometry
PAGE_W_PT, PAGE_H_PT = 612.0, 792.0
MARGIN_L, MARGIN_R = 34.0, 32.0
MARGIN_T, MARGIN_B = 69.0, 74.0
HEADER_DIST, FOOTER_DIST = 28.0, 24.0
COL_W_PT = PAGE_W_PT - MARGIN_L - MARGIN_R          # 546 pt text column
TBL_W_TWIPS = pt_to_twips(COL_W_PT - 2.0)           # official tables stop 2pt short

WATERMARK_RECT = (83.7, 331.9, 447.9, 128.2)        # x, y, w, h
LOGO_RECT = (257.3, 33.4, 100.8, 24.6)
BUTTONS = (
    ("btn_whatsapp.png", 38.5, 722.0, 85.0, 25.0,
     "https://whatsapp.com/channel/0029Van2ybPIt5rqKRSCzK14"),
    ("btn_pw_webapp.png", 265.5, 722.2, 86.2, 25.0,
     "https://smart.link/7wwosivoicgd4"),
    ("btn_telegram.png", 487.8, 722.0, 82.5, 25.0,
     "https://t.me/pw_upsc"),
)

# The same published ads as Daily CA English, with Hindi page furniture.
# Leave room above the Hindi footer buttons (their top edge is 722 pt).
PROMO_ASSETS = os.path.join(ROOT, "assets", "template", "promo")
PROMO_IMAGES = (
    ("promo_srijan.jpeg", fitz.Rect(57.5, 76.3, 560.1, 323.1)),
    ("promo_pyq32.jpeg", fitz.Rect(29.2, 328.7, 582.8, 660.0)),
    ("customer_service_icon.jpeg", fitz.Rect(89.3, 668.0, 136.0, 713.0)),
    ("pw_webapp_line.png", fitz.Rect(145.8, 663.4, 527.8, 715.7)),
)
PROMO_LINKS = (
    (fitz.Rect(299, 663.4, 439, 681.4), "https://smart.link/7wwosivoicgd4"),
    (fitz.Rect(426, 681.4, 518, 699.4), "https://t.me/pw_upsc"),
)

BANNER_H_PT = 38.7
BANNER_ICON_W_PT = 42.8
BANNER_TEXT_W_PT = 501.2
BANNER_ICON_PT = 29.0                               # icon side, centred in the red block
NEWS_BAR_H_PT = 30.0                                # "संक्षेप में समाचार" red bar

PAGENO_BOX = fitz.Rect(289.5, 750.2, 289.5 + 25.5, 750.2 + 25.6)
PAGENO_SIZE = 11.0
FIRST_PAGE_NUMBER = 2
ARIAL_NARROW_BOLD = r"C:\Windows\Fonts\ARIALNB.TTF"

# The official sample carries a faint grey double rule ~16 pt in from the page
# edge on every page (2.0 pt outer + 0.67 pt inner, #939598).  Set to False for
# a border-less page.
DRAW_PAGE_FRAME = True
FRAME_OUTER = fitz.Rect(15.6, 20.5, 596.4, 764.0)
FRAME_INNER = fitz.Rect(18.3, 23.2, 593.7, 761.3)

# ---------------------------------------------------------------- type scale
SZ_TOPIC = 26        # half-points -> 13.0 pt
SZ_HEADING = 22      # 11.0 pt
SZ_GS = 22           # 11.0 pt
SZ_BODY = 23         # 11.5 pt
SZ_NEWSBAR = 26      # 13.0 pt
LINE_PT = 15.0       # measured leading of the official body text

# level (1-based) -> (symbol font, char, half-points, colour)
BULLETS = {
    1: ("Wingdings", "F06C", 14, RED),      # filled circle,  7 pt, red
    2: ("Wingdings", "F06E", 10, INK),      # filled square,  5 pt, black
    3: ("Wingdings 2", "F033", 12, INK),    # tick,           6 pt, black
}
INDENT_STEP_PT = 12.0

# Devanagari must be declared on the complex-script slot too, otherwise Word
# renders the wrong face.  Neither of the two design faces ships with Windows,
# so we walk a fallback chain and report what was actually used.
# The official topic banners are set in Baloo 2 *Bold*. Baloo 2 ships from
# Google Fonts as a single variable font, and Word renders that at its default
# (Regular) instance no matter which named weight is requested - so bold never
# appeared. tools/install_fonts.py cuts real static Regular/Bold instances out
# of it, after which plain "Baloo 2" + bold behaves normally.
DISPLAY_CHAIN = ("Baloo 2", "Nirmala UI", "Mangal", "Arial")
# Body: the published booklet is set in Adobe Devanagari, so that stays first
# and a machine that has it (the typesetting team's does) reproduces the
# sample exactly. It is commercial and cannot ship with the app, so the next
# choice is **Baloo 2** - the face the client asked for by name, and one we
# bundle, so the tool is never at the mercy of what happens to be installed.
# Nirmala UI stays behind it as the always-present Windows safety net.
BODY_CHAIN = ("Adobe Devanagari", "Baloo 2", "Nirmala UI", "Mangal", "Arial")

DEVANAGARI_RE = re.compile(r"[\u0900-\u097F]")
LATIN_RE = re.compile(r"[A-Za-z]")
GS_RE = re.compile(r"^GS[\s\-\u2010-\u2015:_.]*(?:[IVXivx]+|\d+)\b")


# ================================================================ fonts

_font_cache = {}
_font_lock = threading.Lock()


def _installed_font_names():
    """Lower-cased family names Word can actually render on this machine."""
    names = set()
    # 1) Word itself is authoritative (Application.FontNames)
    try:
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        word = None
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            fn = word.FontNames
            names.update(str(fn.Item(i)).strip().lower()
                         for i in range(1, int(fn.Count) + 1))
        finally:
            try:
                if word is not None:
                    word.Quit()
            except Exception:
                pass
            pythoncom.CoUninitialize()
    except Exception:
        pass
    # 2) fall back to the installed-fonts registry
    if not names:
        try:
            import winreg
            key_path = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"
            for root in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
                try:
                    key = winreg.OpenKey(root, key_path)
                except OSError:
                    continue
                i = 0
                while True:
                    try:
                        value_name, _, _ = winreg.EnumValue(key, i)
                    except OSError:
                        break
                    i += 1
                    label = re.sub(r"\s*\((TrueType|OpenType|VGA res)\)\s*$", "",
                                   value_name, flags=re.I)
                    for part in label.split("&"):
                        names.add(part.strip().lower())
        except Exception:
            pass
    return names


def _pick(chain, available):
    """First family of `chain` this machine can render, and whether it is the
    designed one.

    Matching is deliberately tolerant: a family is accepted on an exact name,
    on "Family Style" (how a .ttc or a styled face registers), or ignoring
    spaces - Adobe Creative Cloud has been seen to expose the body face as
    "AdobeDevanagari" rather than "Adobe Devanagari", and missing it means the
    booklet silently comes out in the fallback face, two pages longer than the
    client's own sample. The name that actually matched is what gets returned,
    because Word matches font names case-insensitively but NOT across a
    missing space.
    """
    squashed = {n.replace(" ", ""): n for n in available}
    for family in chain:
        low = family.lower()
        for name in available:
            if name == low or name.startswith(low + " "):
                return family, family == chain[0]
        hit = squashed.get(low.replace(" ", ""))
        if hit:
            return hit, family == chain[0]
    return chain[-1], False


def resolve_fonts(notes=None):
    """(display_font, body_font) with a note about any substitution."""
    with _font_lock:
        if "fonts" not in _font_cache:
            available = _installed_font_names()
            display, display_ok = _pick(DISPLAY_CHAIN, available)
            body, body_ok = _pick(BODY_CHAIN, available)
            if not available:                    # nothing detectable - stay safe
                display, body = "Nirmala UI", "Nirmala UI"
                display_ok = body_ok = False
            _font_cache["fonts"] = (display, body, display_ok, body_ok)
    display, body, display_ok, body_ok = _font_cache["fonts"]
    if notes is not None:
        bits = [f"display '{display}'" + ("" if display_ok else f" (‘{DISPLAY_CHAIN[0]}’ not installed)"),
                f"body '{body}'" + ("" if body_ok else f" (‘{BODY_CHAIN[0]}’ not installed)")]
        notes.append("Fonts used: " + ", ".join(bits) + ".")
        if not body_ok:
            # The published booklet sets its body in Adobe Devanagari 11.5pt.
            # That face is commercial, so the app ships Baloo 2 - the one the
            # client asked for by name - and uses it automatically. Nothing to
            # install for a good result; installing Adobe Devanagari only buys
            # the last step to a byte-identical match.
            notes.append(
                f"Body set in '{body}', which ships with DocPilot. The "
                f"published sample uses '{BODY_CHAIN[0]}' (commercial, part of "
                f"Adobe Creative Cloud): install it on this PC if you want the "
                f"booklet to match the printed copy line for line - DocPilot "
                f"switches to it automatically, no setting to change.")
    return display, body


# ================================================================ xml helpers

def _runs(p_el):
    """Direct run children, including runs nested inside hyperlinks."""
    out = []
    for child in p_el:
        if child.tag == qn('w:r'):
            out.append(child)
        elif child.tag == qn('w:hyperlink'):
            out.extend(c for c in child if c.tag == qn('w:r'))
    return out


def _run_text(r):
    return "".join(t.text or "" for t in r.findall(qn('w:t')))


def _text(el):
    return "".join(t.text or "" for t in el.iter(qn('w:t')))


def _is_bold(r):
    rpr = r.find(qn('w:rPr'))
    if rpr is None:
        return False
    b = rpr.find(qn('w:b'))
    return b is not None and b.get(qn('w:val')) not in ('0', 'false')


def _fully_bold(p_el):
    """True when every run that carries visible text is bold."""
    seen = False
    for r in _runs(p_el):
        if not _run_text(r).strip():
            continue
        seen = True
        if not _is_bold(r):
            return False
    return seen


def _get_rpr(r):
    rpr = r.find(qn('w:rPr'))
    if rpr is None:
        rpr = r.makeelement(qn('w:rPr'), {})
        r.insert(0, rpr)
    return rpr


def _get_ppr(p_el):
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        ppr = p_el.makeelement(qn('w:pPr'), {})
        p_el.insert(0, ppr)
    return ppr


def _get_or_add(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = parent.makeelement(qn(tag), {})
        parent.append(el)
    return el


def _drop(parent, tag):
    for el in parent.findall(qn(tag)):
        parent.remove(el)


def _set_val(parent, tag, val):
    _get_or_add(parent, tag).set(qn('w:val'), val)


def _style_name(p_el, doc):
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        return ""
    ps = ppr.find(qn('w:pStyle'))
    if ps is None:
        return ""
    sid = ps.get(qn('w:val'))
    try:
        node = doc.styles.element.xpath(f'w:style[@w:styleId="{sid}"]/w:name')[0]
        return (node.get(qn('w:val')) or sid).strip().lower()
    except Exception:
        return (sid or "").strip().lower()


def _ilvl(p_el):
    """Source list level, 0-based, or None.  Must be read BEFORE restyling."""
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        return None
    npr = ppr.find(qn('w:numPr'))
    if npr is None:
        return None
    num = npr.find(qn('w:numId'))
    if num is not None and num.get(qn('w:val')) == '0':
        return None
    lvl = npr.find(qn('w:ilvl'))
    try:
        return int(lvl.get(qn('w:val'))) if lvl is not None else 0
    except (TypeError, ValueError):
        return 0


def _style_list_levels(doc):
    """{styleId: ilvl} for paragraph styles that carry their own numbering.

    Typists who reach for Word's built-in 'List Bullet' / 'List Paragraph'
    styles get no w:numPr on the paragraph itself, so the level lives in
    styles.xml.  Cached on the Document.
    """
    cached = getattr(doc, "_hi_style_levels", None)
    if cached is not None:
        return cached
    levels = {}
    try:
        styles = doc.styles.element
        raw = {}
        for style in styles.findall(qn('w:style')):
            sid = style.get(qn('w:styleId'))
            if not sid:
                continue
            based = style.find(qn('w:basedOn'))
            ppr = style.find(qn('w:pPr'))
            npr = ppr.find(qn('w:numPr')) if ppr is not None else None
            lvl = None
            if npr is not None:
                num = npr.find(qn('w:numId'))
                if num is None or num.get(qn('w:val')) != '0':
                    il = npr.find(qn('w:ilvl'))
                    try:
                        lvl = int(il.get(qn('w:val'))) if il is not None else 0
                    except (TypeError, ValueError):
                        lvl = 0
            raw[sid] = (lvl, based.get(qn('w:val')) if based is not None else None)
        for sid in raw:
            seen, cur = set(), sid
            while cur in raw and cur not in seen:
                seen.add(cur)
                lvl, parent = raw[cur]
                if lvl is not None:
                    levels[sid] = lvl
                    break
                cur = parent
    except Exception:
        pass
    try:
        doc._hi_style_levels = levels
    except Exception:
        pass
    return levels


def _list_level(p_el, doc):
    """Source list level, 0-based, or None.  Read BEFORE any restyling."""
    level = _ilvl(p_el)
    if level is not None:
        return level
    ppr = p_el.find(qn('w:pPr'))
    ps = ppr.find(qn('w:pStyle')) if ppr is not None else None
    if ps is None:
        return None
    return _style_list_levels(doc).get(ps.get(qn('w:val')))


def _has_drawing(el):
    return next(el.iter(qn('w:drawing')), None) is not None


def _reset_ppr(p_el, jc="both", left_pt=0.0, hanging_pt=0.0, before=0.0, after=1.0,
               line=LINE_PT, rule="atLeast", keep_next=False, shade=None):
    """Rebuild w:pPr from scratch in schema order.

    Word silently ignores pPr children that appear out of sequence (keepNext
    tacked on after w:jc is the classic way to lose it), and the raw typed file
    brings its own numbering / styles / shading that all have to go.  Only the
    paragraph-mark run properties and a section break are carried over.
    """
    old = p_el.find(qn('w:pPr'))
    rpr = old.find(qn('w:rPr')) if old is not None else None
    sect = old.find(qn('w:sectPr')) if old is not None else None
    if old is not None:
        p_el.remove(old)

    parts = []
    if keep_next:
        parts.append('<w:keepNext/><w:keepLines/>')
    if shade:
        parts.append(f'<w:shd w:val="clear" w:color="auto" w:fill="{shade}"/>')
    spacing = (f'<w:spacing w:before="{pt_to_twips(before)}" w:after="{pt_to_twips(after)}"')
    if line:
        spacing += f' w:line="{pt_to_twips(line)}" w:lineRule="{rule}"'
    parts.append(spacing + '/>')
    ind = f'<w:ind w:left="{pt_to_twips(left_pt)}" w:right="0"'
    ind += (f' w:hanging="{pt_to_twips(hanging_pt)}"' if hanging_pt else ' w:firstLine="0"')
    parts.append(ind + '/>')
    parts.append(f'<w:jc w:val="{jc}"/>')

    ppr = parse_xml(f'<w:pPr {NS}>' + "".join(parts) + '</w:pPr>')
    if rpr is not None:
        ppr.append(rpr)
    if sect is not None:
        ppr.append(sect)
    p_el.insert(0, ppr)
    return ppr


def _shading(ppr_or_tcpr, fill):
    shd = _get_or_add(ppr_or_tcpr, 'w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)


def _paint_mark(p_el, font, size, color=INK, bold=None):
    """The paragraph mark carries its own run properties; left at the typist's
    14 pt it silently inflates every line box."""
    ppr = _get_ppr(p_el)
    rpr = ppr.find(qn('w:rPr'))
    if rpr is None:
        rpr = ppr.makeelement(qn('w:rPr'), {})
        sect = ppr.find(qn('w:sectPr'))
        if sect is not None:
            sect.addprevious(rpr)
        else:
            ppr.append(rpr)
    set_run_fonts(rpr, ascii_font=font, cs_font=font, size_half_pts=size)
    _drop(rpr, 'w:highlight')
    _set_val(rpr, 'w:color', color)
    if bold is not None:
        for tag in ('w:b', 'w:bCs'):
            el = _get_or_add(rpr, tag)
            if bold:
                el.attrib.pop(qn('w:val'), None)
            else:
                el.set(qn('w:val'), '0')


def _paint_runs(p_el, font, cs_font, size, color=INK, bold=None):
    _paint_mark(p_el, font, size, color, bold)
    for r in _runs(p_el):
        rpr = _get_rpr(r)
        set_run_fonts(rpr, ascii_font=font, cs_font=cs_font, size_half_pts=size)
        for tag in ('w:highlight', 'w:shd'):        # typists' marker colours
            _drop(rpr, tag)
        _set_val(rpr, 'w:color', color)
        if bold is not None:
            for tag in ('w:b', 'w:bCs'):
                el = _get_or_add(rpr, tag)
                if bold:
                    el.attrib.pop(qn('w:val'), None)
                else:
                    el.set(qn('w:val'), '0')


def _set_text(p_el, text, font, size, color=INK, bold=False, jc="left"):
    """Replace the paragraph's runs with one clean run (headings / banners only)."""
    for r in list(_runs(p_el)):
        r.getparent().remove(r)
    bold_xml = '<w:b/><w:bCs/>' if bold else ''
    run = parse_xml(
        f'<w:r {NS}><w:rPr>'
        f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}"/>'
        f'{bold_xml}<w:color w:val="{color}"/>'
        f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/>'
        f'</w:rPr><w:t xml:space="preserve"></w:t></w:r>')
    run.find(qn('w:t')).text = text
    p_el.append(run)
    ppr = _get_ppr(p_el)
    _set_val(ppr, 'w:jc', jc)
    _paint_mark(p_el, font, size, color, bold)


# ================================================================ drawings

def _inline_image_xml(rid, w_pt, h_pt, name, did):
    return f'''<w:drawing {NS}>
  <wp:inline distT="0" distB="0" distL="0" distR="0">
    <wp:extent cx="{pt_to_emu(w_pt)}" cy="{pt_to_emu(h_pt)}"/>
    <wp:effectExtent l="0" t="0" r="0" b="0"/>
    <wp:docPr id="{did}" name="{name}"/>
    <wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
    <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
      <pic:pic>
        <pic:nvPicPr><pic:cNvPr id="{did}" name="{name}"/><pic:cNvPicPr/></pic:nvPicPr>
        <pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
        <pic:spPr><a:xfrm><a:off x="0" y="0"/>
          <a:ext cx="{pt_to_emu(w_pt)}" cy="{pt_to_emu(h_pt)}"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>
      </pic:pic>
    </a:graphicData></a:graphic>
  </wp:inline>
</w:drawing>'''


_did = [7000]


def _next_did():
    _did[0] += 1
    return _did[0]


def _add_inline_image(p_el, part, path, w_pt, h_pt, name="img"):
    rid, _ = part.get_or_add_image(path)
    run = parse_xml(f'<w:r {NS}></w:r>')
    run.append(parse_xml(_inline_image_xml(rid, w_pt, h_pt, name, _next_did())))
    p_el.append(run)
    return run


def _drawing_size(holder):
    ext = holder.find(qn('wp:extent'))
    if ext is None:
        return None
    try:
        return int(ext.get('cx')), int(ext.get('cy'))
    except (TypeError, ValueError):
        return None


def _resize_drawing(holder, w_emu, h_emu):
    ext = holder.find(qn('wp:extent'))
    if ext is not None:
        ext.set('cx', str(w_emu))
        ext.set('cy', str(h_emu))
    for a_ext in holder.iter(qn('a:ext')):
        a_ext.set('cx', str(w_emu))
        a_ext.set('cy', str(h_emu))


def _cap_width(holder, max_w_pt):
    size = _drawing_size(holder)
    if size is None:
        return
    cx, cy = size
    max_emu = pt_to_emu(max_w_pt)
    if cx <= max_emu or cx <= 0:
        return
    factor = max_emu / cx
    _resize_drawing(holder, int(cx * factor), int(cy * factor))


def _anchor_xml(w_emu, h_emu, name, did):
    return f'''<wp:anchor {NS} distT="0" distB="0" distL="127000" distR="0"
        simplePos="0" relativeHeight="{251660000 + did}" behindDoc="0" locked="0"
        layoutInCell="1" allowOverlap="1">
  <wp:simplePos x="0" y="0"/>
  <wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH>
  <wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>
  <wp:extent cx="{w_emu}" cy="{h_emu}"/>
  <wp:effectExtent l="0" t="0" r="0" b="0"/>
  <wp:wrapSquare wrapText="bothSides"/>
  <wp:docPr id="{did}" name="{name}"/>
  <wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
</wp:anchor>'''


def _dock_right(anchor):
    """Pin an existing floating picture to the right text margin."""
    pos_h = anchor.find(qn('wp:positionH'))
    if pos_h is None:
        return
    pos_h.set('relativeFrom', 'margin')
    for child in list(pos_h):
        pos_h.remove(child)
    align = pos_h.makeelement(qn('wp:align'), {})
    align.text = 'right'
    pos_h.append(align)


def _float_body_images(body, max_w_pt):
    """Picture-only paragraphs become a right-docked, text-wrapped float.

    The official layout runs the body text down the left of every illustration;
    a stand-alone inline picture would instead eat a full-width band.  The
    drawing is moved into the following paragraph and the now-empty carrier
    paragraph is dropped (never the picture itself).
    """
    floated = 0
    for p_el in [el for el in body if el.tag == qn('w:p')]:
        if _text(p_el).strip():
            continue
        inlines = list(p_el.iter(qn('wp:inline')))
        if len(inlines) != 1:
            continue
        holder = inlines[0]
        size = _drawing_size(holder)
        if size is None:
            continue
        _cap_width(holder, max_w_pt)
        cx, cy = _drawing_size(holder)
        _reset_ppr(p_el, jc='center', before=2, after=2, line=None)
        nxt = p_el.getnext()
        while nxt is not None and nxt.tag not in (qn('w:p'), qn('w:tbl')):
            nxt = nxt.getnext()
        if nxt is None or nxt.tag != qn('w:p') or not _text(nxt).strip():
            continue                                  # keep it inline & centred
        after_nxt = nxt.getnext()
        while after_nxt is not None and after_nxt.tag not in (qn('w:p'), qn('w:tbl')):
            after_nxt = after_nxt.getnext()
        if after_nxt is not None and after_nxt.tag == qn('w:tbl'):
            # a float here would push the (unsplittable) table onto the next
            # page and leave a half-empty one behind
            continue
        run = holder.getparent()
        while run is not None and run.tag != qn('w:r'):
            run = run.getparent()
        if run is None:
            continue
        anchor = parse_xml(_anchor_xml(cx, cy, "content_image", _next_did()))
        for child in list(holder):
            if child.tag in (qn('wp:extent'), qn('wp:effectExtent'),
                             qn('wp:docPr'), qn('wp:cNvGraphicFramePr')):
                continue
            anchor.append(child)                       # moves a:graphic across
        holder.getparent().replace(holder, anchor)
        ppr = nxt.find(qn('w:pPr'))
        nxt.insert(1 if ppr is not None else 0, run)
        p_el.getparent().remove(p_el)
        floated += 1
    # already-floating pictures: cap them and dock them right as well
    for anchor in body.iter(qn('wp:anchor')):
        if anchor.find(qn('wp:wrapNone')) is not None:
            continue                                   # template art
        _cap_width(anchor, max_w_pt)
        _dock_right(anchor)
    return floated


def _cap_cell_images(tbl_el, widths_pt):
    for tr in tbl_el.findall(qn('w:tr')):
        for idx, tc in enumerate(tr.findall(qn('w:tc'))):
            limit = widths_pt[idx] - 10 if idx < len(widths_pt) else COL_W_PT
            for tag in ('wp:inline', 'wp:anchor'):
                for holder in tc.iter(qn(tag)):
                    _cap_width(holder, max(40.0, limit))


# ================================================================ paragraphs

def _style_body(p_el, fonts, justify=True):
    display, body = fonts
    _reset_ppr(p_el, jc='both' if justify else 'left', before=0, after=1)
    _paint_runs(p_el, body, body, SZ_BODY)


def _style_heading(p_el, fonts, keep_next=True):
    display, body = fonts
    # keepNext inside a table cell chains whole rows together and stops Word
    # breaking a tall table across pages - callers in cells pass False
    _reset_ppr(p_el, jc='left', before=4, after=1, keep_next=keep_next)
    _paint_runs(p_el, display, display, SZ_HEADING, bold=True)


def _style_gs_line(p_el, fonts):
    display, body = fonts
    # keepNext: the banner keeps to this line and this line to the first
    # heading, so a topic head never sits alone at a page foot (seen once
    # topics flowed again, 2026-10-12: banner + GS line at the foot of p4,
    # the topic's text overleaf)
    _reset_ppr(p_el, jc='right', before=3, after=3, keep_next=True)
    _set_text(p_el, _text(p_el).strip(), body, SZ_GS, INK, bold=False, jc="right")


def _style_list_item(p_el, level, fonts):
    """Literal Wingdings bullet + tab, hanging indent.  Word's list engine is
    not reliable for pinning a specific glyph/size/colour, so we write it."""
    display, body = fonts
    _reset_ppr(p_el, jc='both', left_pt=INDENT_STEP_PT * level,
               hanging_pt=INDENT_STEP_PT, before=0, after=1)
    _paint_runs(p_el, body, body, SZ_BODY)

    sym_font, char, size, color = BULLETS.get(min(level, max(BULLETS)), BULLETS[3])
    bullet = parse_xml(
        f'<w:r {NS}><w:rPr>'
        f'<w:rFonts w:ascii="{sym_font}" w:hAnsi="{sym_font}" w:cs="{sym_font}" w:hint="default"/>'
        f'<w:b w:val="0"/><w:bCs w:val="0"/>'
        f'<w:color w:val="{color}"/><w:sz w:val="{size}"/><w:szCs w:val="{size}"/>'
        f'</w:rPr><w:sym w:font="{sym_font}" w:char="{char}"/></w:r>')
    tab = parse_xml(
        f'<w:r {NS}><w:rPr>'
        f'<w:rFonts w:ascii="{body}" w:hAnsi="{body}" w:cs="{body}"/>'
        f'<w:sz w:val="{SZ_BODY}"/><w:szCs w:val="{SZ_BODY}"/>'
        f'</w:rPr><w:tab/></w:r>')
    at = 1 if p_el.find(qn('w:pPr')) is not None else 0
    p_el.insert(at, tab)
    p_el.insert(at, bullet)


def _style_cell_paragraph(p_el, doc, fonts):
    level = _list_level(p_el, doc)
    if level is not None:
        _style_list_item(p_el, level + 1, fonts)
    elif _text(p_el).strip() and _fully_bold(p_el):
        _style_heading(p_el, fonts, keep_next=False)
    else:
        _style_body(p_el, fonts)


# ================================================================ banners

def _bar_table(widths_pt, fills, height_pt, before_pt=0.0):
    """A borderless single-row table used for the coloured banner bars.

    Paragraph shading only ever covers the text's own line box, so a bar of an
    exact height has to be a table row - the topic banner and the news bar are
    both built this way.
    """
    widths = [pt_to_twips(w) for w in widths_pt]
    cells = ""
    for width, fill in zip(widths, fills):
        cells += (
            f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>'
            f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>'
            f'<w:tcMar><w:left w:w="{pt_to_twips(4.5)}" w:type="dxa"/>'
            f'<w:right w:w="{pt_to_twips(4.5)}" w:type="dxa"/></w:tcMar>'
            f'<w:vAlign w:val="center"/></w:tcPr>'
            f'<w:p><w:pPr><w:keepNext/><w:keepLines/>'
            f'<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
            f'<w:ind w:left="0" w:right="0" w:firstLine="0"/>'
            f'<w:jc w:val="center"/></w:pPr></w:p></w:tc>')
    spacing_row = (f'<w:tr><w:trPr>'
                   f'<w:trHeight w:val="{pt_to_twips(before_pt)}" w:hRule="exact"/>'
                   f'</w:trPr><w:tc><w:tcPr>'
                   f'<w:tcW w:w="{sum(widths)}" w:type="dxa"/>'
                   f'<w:gridSpan w:val="{len(widths)}"/></w:tcPr>'
                   f'<w:p><w:pPr><w:spacing w:before="0" w:after="0" '
                   f'w:line="20" w:lineRule="exact"/></w:pPr></w:p>'
                   f'</w:tc></w:tr>') if before_pt else ""
    return parse_xml(
        f'<w:tbl {NS}>'
        f'<w:tblPr>'
        f'<w:tblW w:w="{sum(widths)}" w:type="dxa"/>'
        f'<w:tblLayout w:type="fixed"/>'
        f'<w:tblInd w:w="0" w:type="dxa"/>'
        f'<w:tblCellMar>'
        f'<w:top w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/>'
        f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/>'
        f'</w:tblCellMar>'
        f'<w:tblBorders>' + "".join(
            f'<w:{e} w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            for e in ("top", "left", "bottom", "right", "insideH", "insideV")) +
        f'</w:tblBorders>'
        f'</w:tblPr>'
        f'<w:tblGrid>' + "".join(f'<w:gridCol w:w="{w}"/>' for w in widths) +
        f'</w:tblGrid>'
        + spacing_row +
        f'<w:tr><w:trPr><w:cantSplit/>'
        f'<w:trHeight w:val="{pt_to_twips(height_pt)}" w:hRule="atLeast"/>'
        f'</w:trPr>{cells}</w:tr>'
        f'</w:tbl>')


def _swap_for_table(p_el, tbl):
    p_el.addprevious(tbl)
    p_el.getparent().remove(p_el)
    return tbl


def _banner_table(p_el, doc, fonts):
    """Topic title -> red icon block + peach banner bar (1 row, 2 cells)."""
    title = " ".join(_text(p_el).split())
    tbl = _bar_table([BANNER_ICON_W_PT, BANNER_TEXT_W_PT], [RED, PEACH],
                     BANNER_H_PT, before_pt=2.0)
    cells = tbl.findall('.//' + qn('w:tc'))[-2:]
    icon_path = os.path.join(ASSETS, "topic_news_icon.png")
    if os.path.exists(icon_path):
        _add_inline_image(cells[0].find(qn('w:p')), doc.part, icon_path,
                          BANNER_ICON_PT, BANNER_ICON_PT, name="topic_icon")
    _set_text(cells[1].find(qn('w:p')), title, fonts[0], SZ_TOPIC, RED,
              bold=True, jc="left")
    return _swap_for_table(p_el, tbl)


def _news_bar_table(p_el, doc, fonts):
    """'संक्षेप में समाचार' -> full-width red bar, white centred display text."""
    tbl = _bar_table([BANNER_ICON_W_PT + BANNER_TEXT_W_PT], [RED],
                     NEWS_BAR_H_PT, before_pt=6.0)
    cell = tbl.findall('.//' + qn('w:tc'))[-1]
    _set_text(cell.find(qn('w:p')), " ".join(_text(p_el).split()),
              fonts[0], SZ_NEWSBAR, WHITE, bold=True, jc="center")
    return _swap_for_table(p_el, tbl)


# ================================================================ tables

def _table_borders(tbl_el, color, size_eighths, inside=True):
    tblpr = tbl_el.find(qn('w:tblPr'))
    if tblpr is None:
        tblpr = tbl_el.makeelement(qn('w:tblPr'), {})
        tbl_el.insert(0, tblpr)
    _drop(tblpr, 'w:tblBorders')
    edges = ["top", "left", "bottom", "right"]
    if inside:
        edges += ["insideH", "insideV"]
    tblpr.append(parse_xml(
        f'<w:tblBorders {NS}>' + "".join(
            f'<w:{e} w:val="single" w:sz="{size_eighths}" w:space="0" w:color="{color}"/>'
            for e in edges) + '</w:tblBorders>'))
    return tblpr


def _fit_table(tbl_el):
    """Scale the source grid onto the Hindi text column; returns widths in pt."""
    tblpr = tbl_el.find(qn('w:tblPr'))
    if tblpr is None:
        tblpr = tbl_el.makeelement(qn('w:tblPr'), {})
        tbl_el.insert(0, tblpr)
    for tag in ('w:tblW', 'w:tblLayout', 'w:tblInd'):
        _drop(tblpr, tag)
    tblpr.append(parse_xml(f'<w:tblW {NS} w:w="{TBL_W_TWIPS}" w:type="dxa"/>'))
    tblpr.append(parse_xml(f'<w:tblLayout {NS} w:type="fixed"/>'))
    tblpr.append(parse_xml(f'<w:tblInd {NS} w:w="0" w:type="dxa"/>'))
    for cant in list(tbl_el.iter(qn('w:cantSplit'))):
        cant.getparent().remove(cant)

    rows = tbl_el.findall(qn('w:tr'))
    ncols = max((len(tr.findall(qn('w:tc'))) for tr in rows), default=0)
    if not ncols:
        return []
    grid = tbl_el.find(qn('w:tblGrid'))
    src = []
    if grid is not None:
        for col in grid.findall(qn('w:gridCol')):
            try:
                src.append(int(col.get(qn('w:w'))))
            except (TypeError, ValueError):
                src.append(0)
    src = [w for w in src if w > 0][:ncols]
    if len(src) != ncols or sum(src) <= 0:
        src = [1] * ncols
    total = sum(src)
    widths = [int(round(TBL_W_TWIPS * w / total)) for w in src]
    widths[-1] += TBL_W_TWIPS - sum(widths)

    if grid is not None:
        tbl_el.remove(grid)
    grid = parse_xml(f'<w:tblGrid {NS}>' +
                     "".join(f'<w:gridCol w:w="{w}"/>' for w in widths) +
                     '</w:tblGrid>')
    tblpr.addnext(grid)
    for tr in rows:
        for idx, tc in enumerate(tr.findall(qn('w:tc'))):
            tcpr = tc.find(qn('w:tcPr'))
            if tcpr is None:
                tcpr = tc.makeelement(qn('w:tcPr'), {})
                tc.insert(0, tcpr)
            _drop(tcpr, 'w:tcW')
            span = tcpr.find(qn('w:gridSpan'))
            n = 1
            if span is not None:
                try:
                    n = max(1, int(span.get(qn('w:val'))))
                except (TypeError, ValueError):
                    n = 1
            width = sum(widths[idx:idx + n]) or widths[min(idx, len(widths) - 1)]
            tcpr.insert(0, parse_xml(f'<w:tcW {NS} w:w="{width}" w:type="dxa"/>'))
    return [w / 20.0 for w in widths]


def _cell_paragraphs(tc):
    return [p for p in tc if p.tag == qn('w:p')]


def _style_infobox(tbl_el, doc, fonts, widths):
    """1x1 table -> peach panel with a thin red rule."""
    _table_borders(tbl_el, RED, 4, inside=False)
    for tc in tbl_el.iter(qn('w:tc')):
        tcpr = tc.find(qn('w:tcPr'))
        if tcpr is not None:
            _shading(tcpr, PEACH_SOFT)
        for p in _cell_paragraphs(tc):
            _style_cell_paragraph(p, doc, fonts)


def _style_news_table(tbl_el, doc, fonts, widths):
    """N x 2 'News in shorts' grid: red rules, bold display-font left column."""
    _table_borders(tbl_el, RED, 4, inside=True)
    for tr in tbl_el.findall(qn('w:tr')):
        tcs = tr.findall(qn('w:tc'))
        if len(tcs) != 2:
            for tc in tcs:
                for p in _cell_paragraphs(tc):
                    _style_cell_paragraph(p, doc, fonts)
            continue
        left, right = tcs
        for tcpr_owner in (left, right):
            tcpr = tcpr_owner.find(qn('w:tcPr'))
            if tcpr is not None:
                _drop(tcpr, 'w:shd')
                _get_or_add(tcpr, 'w:vAlign').set(qn('w:val'), 'top')
        for p in _cell_paragraphs(left):
            _reset_ppr(p, jc='center', before=2, after=2)
            _paint_runs(p, fonts[0], fonts[0], SZ_HEADING, INK, bold=True)
        for p in _cell_paragraphs(right):
            _style_cell_paragraph(p, doc, fonts)


def _style_grid_table(tbl_el, doc, fonts, widths):
    """>=3 column table: red header row, peach body, white rules."""
    _table_borders(tbl_el, WHITE, 8, inside=True)
    rows = tbl_el.findall(qn('w:tr'))
    for ri, tr in enumerate(rows):
        header = (ri == 0)
        trpr = tr.find(qn('w:trPr'))
        if header:
            if trpr is None:
                trpr = tr.makeelement(qn('w:trPr'), {})
                tr.insert(0, trpr)
            _get_or_add(trpr, 'w:tblHeader')
        for tc in tr.findall(qn('w:tc')):
            tcpr = tc.find(qn('w:tcPr'))
            if tcpr is not None:
                _shading(tcpr, RED if header else PEACH_SOFT)
                _get_or_add(tcpr, 'w:vAlign').set(qn('w:val'), 'top')
            for p in _cell_paragraphs(tc):
                if header:
                    # keepNext pulls the header row down to the first body row
                    # instead of stranding it alone at the foot of a page
                    _reset_ppr(p, jc='left', before=1, after=1, keep_next=True)
                    _paint_runs(p, fonts[1], fonts[1], SZ_BODY, WHITE, bold=True)
                else:
                    _style_cell_paragraph(p, doc, fonts)


def _style_table(tbl_el, doc, fonts):
    # heading pStyles inside cells drag keepNext chains that stop Word from
    # breaking tall rows across pages
    for ps in list(tbl_el.iter(qn('w:pStyle'))):
        ps.getparent().remove(ps)
    # a Google Docs export pins every row to a minimum height (trHeight
    # atLeast - up to 1650tw, over an inch, on the 17-Aug challenges grid), so
    # a nine-row table filled a page on its own against the published
    # sample's natural rows. Drop the minima; the combo pipeline already does
    # (combo_hindi_layout._content_tables). The banner tables built below set
    # their own heights and never pass through here.
    for tr in tbl_el.findall(qn('w:tr')):
        trpr = tr.find(qn('w:trPr'))
        if trpr is not None:
            for height in trpr.findall(qn('w:trHeight')):
                trpr.remove(height)
    widths = _fit_table(tbl_el)
    rows = tbl_el.findall(qn('w:tr'))
    ncols = len(rows[0].findall(qn('w:tc'))) if rows else 0
    if len(rows) == 1 and ncols == 1:
        _style_infobox(tbl_el, doc, fonts, widths)
    elif ncols == 2:
        _style_news_table(tbl_el, doc, fonts, widths)
    else:
        _style_grid_table(tbl_el, doc, fonts, widths)
    _cap_cell_images(tbl_el, widths)


# ================================================================ page template

def _apply_page_template(doc):
    for i, section in enumerate(doc.sections):
        section.page_width = Pt(PAGE_W_PT)
        section.page_height = Pt(PAGE_H_PT)
        section.left_margin = Pt(MARGIN_L)
        section.right_margin = Pt(MARGIN_R)
        section.top_margin = Pt(MARGIN_T)
        section.bottom_margin = Pt(MARGIN_B)
        section.header_distance = Pt(HEADER_DIST)
        section.footer_distance = Pt(FOOTER_DIST)
        section.different_first_page_header_footer = False
        sect_pr = section._sectPr
        _drop(sect_pr, 'w:pgBorders')            # the Hindi design has none
        _drop(sect_pr, 'w:titlePg')
        if i > 0:
            section.header.is_linked_to_previous = True
            section.footer.is_linked_to_previous = True

    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    _reset_ppr(hp._p, jc="left", before=0, after=0, line=1, rule="exact")
    wm = os.path.join(ASSETS, "watermark_center.png")
    if os.path.exists(wm):
        x, y, w, h = WATERMARK_RECT
        add_anchored_image(hp, header.part, wm, x, y, w, h,
                           behind=True, name="hi_watermark")
    logo = os.path.join(ASSETS, "header_logo.png")
    if os.path.exists(logo):
        x, y, w, h = LOGO_RECT
        add_anchored_image(hp, header.part, logo, x, y, w, h, name="hi_logo")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    _reset_ppr(fp._p, jc="left", before=0, after=0, line=1, rule="exact")
    for name, x, y, w, h, link in BUTTONS:
        path = os.path.join(ASSETS, name)
        if os.path.exists(path):
            add_anchored_image(fp, footer.part, path, x, y, w, h,
                               hyperlink=link, name=os.path.splitext(name)[0])


_REVISION_TAGS = ('w:tblGridChange', 'w:tblPrChange', 'w:trPrChange', 'w:tcPrChange',
                  'w:pPrChange', 'w:rPrChange', 'w:sectPrChange')


def _cleanup_source(body):
    """Strip Google-Docs export artefacts that break Word pagination."""
    changed = True
    while changed:                                   # flatten content controls
        changed = False
        for sdt in list(body.iter(qn('w:sdt'))):
            parent = sdt.getparent()
            if parent is None:
                continue
            content = sdt.find(qn('w:sdtContent'))
            if content is not None:
                for child in list(content):
                    sdt.addprevious(child)
            parent.remove(sdt)
            changed = True
    for tag in _REVISION_TAGS + ('w:pageBreakBefore', 'w:proofErr'):
        for el in list(body.iter(qn(tag))):
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)
    for ind in body.iter(qn('w:tblInd')):
        ind.set(qn('w:w'), '0')
        ind.set(qn('w:type'), 'dxa')


# ================================================================ restyle

def _drop_tab_artifacts(body, doc):
    """Google Docs tabs front matter. The scan itself lives in
    `ooxml.drop_google_tab_frontmatter` so the English/Xylem styler and this
    one cannot drift apart - the same raw file can reach either of them."""
    return drop_google_tab_frontmatter(body, lambda el: _style_name(el, doc))


def restyle_docx(in_path, out_path, fonts=None, notes=None):
    """Raw typed Hindi CA docx -> branded Hindi docx.  Formatting only."""
    notes = notes if notes is not None else []
    fonts = fonts or resolve_fonts()
    doc = Document(in_path)
    _apply_page_template(doc)
    body = doc.element.body
    _cleanup_source(body)
    tabs, tab_breaks = _drop_tab_artifacts(body, doc)
    if tabs or tab_breaks:
        notes.append(f"Google Docs tab front matter removed: {tabs} 'Tab N' "
                     f"paragraph(s), {tab_breaks} page break(s) - they were "
                     f"leaving a near-empty page before the first topic.")

    topics = tables = glosses = dropped = 0
    awaiting = 0                      # paragraphs still eligible as gloss / GS

    for el in [e for e in body if e.tag in (qn('w:p'), qn('w:tbl'))]:
        if el.tag == qn('w:tbl'):
            _style_table(el, doc, fonts)
            tables += 1
            awaiting = 0
            continue

        text = _text(el).strip()
        style = _style_name(el, doc)
        level = _list_level(el, doc)                   # snapshot BEFORE editing

        if not text and not _has_drawing(el):
            el.getparent().remove(el)
            dropped += 1
            continue

        if style == "heading 2":
            _banner_table(el, doc, fonts)
            topics += 1
            awaiting = 2
            continue

        if style == "heading 1":
            _news_bar_table(el, doc, fonts)
            awaiting = 0
            continue

        if awaiting:
            awaiting -= 1
            # the Latin gloss under a topic title is not part of the design
            if not DEVANAGARI_RE.search(text) and LATIN_RE.search(text) \
                    and not _has_drawing(el):
                el.getparent().remove(el)
                glosses += 1
                awaiting = max(awaiting, 1)
                continue
            if GS_RE.match(text):
                _style_gs_line(el, fonts)
                awaiting = 0
                continue
            awaiting = 0

        if level is not None:
            _style_list_item(el, level + 1, fonts)
        elif text and _fully_bold(el):
            _style_heading(el, fonts)
        else:
            _style_body(el, fonts)

    floated = _float_body_images(body, COL_W_PT * 0.40)

    # Word needs a paragraph after a trailing table
    blocks = [e for e in body if e.tag in (qn('w:p'), qn('w:tbl'))]
    if blocks and blocks[-1].tag == qn('w:tbl'):
        tail = parse_xml(f'<w:p {NS}><w:pPr>'
                         f'<w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>'
                         f'</w:pPr></w:p>')
        blocks[-1].addnext(tail)

    notes.append(f"{topics} topic banner(s), {tables} table(s) restyled; "
                 f"{glosses} English gloss line(s) dropped, "
                 f"{dropped} blank paragraph(s) removed, "
                 f"{floated} illustration(s) docked right.")
    if not topics:
        notes.append("No 'Heading 2' topic titles found - the input may not be "
                     "a Daily Current Affairs file.")
    doc.save(out_path)
    return out_path


# ================================================================ pdf branding

_pageno_font = [None]


def _get_pageno_font():
    if _pageno_font[0] is None:
        try:
            _pageno_font[0] = fitz.Font(fontfile=ARIAL_NARROW_BOLD)
        except Exception:
            _pageno_font[0] = fitz.Font("hebo")
    return _pageno_font[0]


def _hex_rgb(value):
    return (int(value[0:2], 16) / 255, int(value[2:4], 16) / 255, int(value[4:6], 16) / 255)


def _stamp_page(page, number):
    if DRAW_PAGE_FRAME:
        grey = _hex_rgb(FRAME_GREY)
        page.draw_rect(FRAME_OUTER, color=grey, width=2.0)
        page.draw_rect(FRAME_INNER, color=grey, width=0.67)
    page.draw_rect(PAGENO_BOX, color=None, fill=_hex_rgb(INK), radius=0.25)
    label = str(number)
    font = _get_pageno_font()
    width = font.text_length(label, fontsize=PAGENO_SIZE)
    baseline = fitz.Point(
        (PAGENO_BOX.x0 + PAGENO_BOX.x1) / 2 - width / 2,
        (PAGENO_BOX.y0 + PAGENO_BOX.y1) / 2
        + (font.ascender + font.descender) / 2 * PAGENO_SIZE)
    if font.name == "Helvetica-Bold":
        page.insert_text(baseline, label, fontname="hebo",
                         fontsize=PAGENO_SIZE, color=(1, 1, 1))
    else:
        page.insert_text(baseline, label, fontfile=ARIAL_NARROW_BOLD,
                         fontname="ArialNarrowBold", fontsize=PAGENO_SIZE,
                         color=(1, 1, 1))


def _build_promo_page(doc, number):
    """Append the dedicated ad page; its presence never depends on body space."""
    page = doc.new_page(width=PAGE_W_PT, height=PAGE_H_PT)
    for name, (x, y, width, height) in (
            ("watermark_center.png", WATERMARK_RECT),
            ("header_logo.png", LOGO_RECT)):
        page.insert_image(fitz.Rect(x, y, x + width, y + height),
                          filename=os.path.join(ASSETS, name))
    for name, rect in PROMO_IMAGES:
        page.insert_image(rect, filename=os.path.join(PROMO_ASSETS, name))
    for name, x, y, width, height, uri in BUTTONS:
        rect = fitz.Rect(x, y, x + width, y + height)
        page.insert_image(rect, filename=os.path.join(ASSETS, name))
        page.insert_link({"kind": fitz.LINK_URI, "from": rect, "uri": uri})
    for rect, uri in PROMO_LINKS:
        page.insert_link({"kind": fitz.LINK_URI, "from": rect, "uri": uri})
    _stamp_page(page, number)
    return page


def finalize_pdf(in_pdf, out_pdf, first_number=FIRST_PAGE_NUMBER):
    """Brand the content and append its numbered closing advertisement page."""
    doc = fitz.open(in_pdf)
    number = first_number
    for page in doc:
        _stamp_page(page, number)
        number += 1
    _build_promo_page(doc, number)
    doc.save(out_pdf, garbage=3, deflate=True)
    doc.close()
    return out_pdf


# ================================================================ pipeline

def output_name(input_stem, variant="Hindi"):
    return f"{input_stem}_CA Analysis_UPSC_{variant}.pdf"


def process(input_path, workdir, variant="Hindi",
            progress=lambda stage, pct: None, cancelled=lambda: False, **kw):
    """Raw typed Hindi CA docx -> branded PDF.

    Returns ``(output_pdf_path, notes_list)``.
    """
    os.makedirs(workdir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(input_path))[0]
    notes = []

    progress("Checking Devanagari fonts", 8)
    fonts = resolve_fonts(notes)
    if cancelled():
        raise InterruptedError("stopped")

    progress("Styling document (Hindi)", 25)
    styled = os.path.join(workdir, f"{stem}_styled_hi.docx")
    restyle_docx(input_path, styled, fonts, notes)
    if cancelled():
        raise InterruptedError("stopped")

    # measured pagination (client mark-up 2026-08-26, same as English): fill
    # page bottoms, pull small tails back, forward pushed float anchors. The
    # pass converts through Word while it measures, so its last export
    # already matches the final docx state.
    progress("Fitting pages", 40)
    raw_pdf = fit_document(styled, workdir, spec=HINDI_SPEC,
                           progress=progress, notes=notes)
    if cancelled():
        raise InterruptedError("stopped")

    if not os.path.exists(raw_pdf):
        progress("Converting to PDF (Word)", 60)
        raw_pdf = os.path.join(workdir, f"{stem}_raw_hi.pdf")
        docx_to_pdf(styled, raw_pdf)
    if cancelled():
        raise InterruptedError("stopped")

    progress("Branding pages & promo", 88)
    final_pdf = os.path.join(workdir, output_name(stem, variant))
    finalize_pdf(raw_pdf, final_pdf, first_number=FIRST_PAGE_NUMBER)

    # editable Word companion, named after the PDF so callers can find it by
    # convention (page numbers and the promo page are stamped into the PDF
    # only - the .docx is the branded, fitted source)
    companion = os.path.splitext(final_pdf)[0] + ".docx"
    try:
        shutil.copy2(styled, companion)
    except OSError:
        pass

    # front page (2026-08-31): a separate stage after every body deliverable
    # exists - it prepends a page to the PDF and the companion and never
    # touches the body; a failure keeps the booklet as it is
    try:
        from .cover_hindi import attach_cover
        attach_cover(input_path, styled, final_pdf, companion, workdir,
                     first_number=FIRST_PAGE_NUMBER,
                     date_in=kw.get("cover_date"), index_in=kw.get("cover_index"),
                     shorts_in=kw.get("cover_shorts"),
                     progress=progress, cancelled=cancelled, notes=notes)
    except InterruptedError:
        raise
    except Exception as exc:
        notes.append(f"No cover page ({type(exc).__name__}: {exc}); the booklet "
                     f"is unchanged.")

    progress("Done", 100)
    return final_pdf, notes
