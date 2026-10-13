"""Regression checks for reusable combined-Hindi layout (no Word required)."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from backend.pipeline import combo_hindi_layout as L, dea_common as C, daily_ca_hindi as H


class CombinedLayoutTests(unittest.TestCase):
    def test_input_row_heights_and_cell_borders_do_not_override_template(self):
        d = Document()
        t = d.add_table(rows=3, cols=3)
        t.cell(0, 0).text = 'Heading'
        t.cell(1, 1).text = 'A different article with several lines of text.'
        before = L._text(t._tbl)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(parse_xml(f'<w:trHeight {C.W_NS} w:val="1600" w:hRule="atLeast"/>'))
        t.cell(1, 1)._tc.get_or_add_tcPr().append(parse_xml(f'<w:tcBorders {C.W_NS}><w:top w:val="single" w:color="C4C7C5"/></w:tcBorders>'))
        L._content_tables(d.element.body)
        self.assertEqual(before, L._text(t._tbl))
        self.assertFalse(list(t._tbl.iter(qn('w:trHeight'))))
        self.assertFalse(list(t._tbl.iter(qn('w:tcBorders'))))
        self.assertTrue(all(e.get(qn('w:w')) == '0' for tc in t._tbl.iter(qn('w:tcMar'))
                            for e in tc if e.tag in (qn('w:top'), qn('w:bottom'))))

    def test_measured_banner_height_survives_table_cleanup(self):
        d = Document()
        p = d.add_paragraph('Any new topic title')._p
        t = H._banner_table(p, d, ('Baloo 2', 'Baloo 2'))
        before = [dict(h.attrib) for h in t.iter(qn('w:trHeight'))]
        L._content_tables(d.element.body)
        self.assertEqual(before, [dict(h.attrib) for h in t.iter(qn('w:trHeight'))])

    def test_editorial_end_marker_precedes_next_day_spacer(self):
        with tempfile.TemporaryDirectory() as wd:
            path = str(Path(wd) / 'variant.docx')
            d = Document()
            H._banner_table(d.add_paragraph('Changed first topic')._p, d, ('Baloo 2', 'Baloo 2'))
            ending = d.add_paragraph('The last paragraph of the first day.')._p
            d.add_paragraph('')
            H._banner_table(d.add_paragraph('Changed next day topic')._p, d, ('Baloo 2', 'Baloo 2'))
            d.add_paragraph('Second day content.')
            before = L._text(d.element.body)
            d.save(path)
            with patch.object(L, '_furniture'):
                L.normalize(path, wd, ('Baloo 2', 'Baloo 2'), editorial=True,
                            topics=['Changed first topic', 'Changed next day topic'], day_starts=[0, 1])
            result = Document(path)
            bm = next(result.element.body.iter(qn('w:bookmarkStart')))
            self.assertEqual(L._text(bm.getparent()), 'The last paragraph of the first day.')
            self.assertEqual(L._text(result.element.body), before)
            self.assertEqual(bm.get(qn('w:name')), L.PREFIX + 'day0')

    def test_word_ad_uses_a_body_anchor_after_a_table(self):
        from backend.pipeline import dea_hindi as D
        with tempfile.TemporaryDirectory() as wd:
            path = str(Path(wd) / 'table_end.docx')
            d = Document()
            table = d.add_table(rows=1, cols=2)
            table.cell(0, 1).text = 'The last news item.'
            L._mark(table.cell(0, 1).paragraphs[0]._p, 'ca', 27001)
            d.save(path)
            L.add_promos(path, {'ca': (D.SRIJAN, (133, 440, 490, 667), 3)})
            result = Document(path)
            anchor = next(result.element.body.iter(qn('wp:anchor')))
            p = next(a for a in anchor.iterancestors() if a.tag == qn('w:p'))
            self.assertEqual(p.getparent(), result.element.body)
            self.assertEqual(L._text(result.tables[0]._tbl), 'The last news item.')


if __name__ == '__main__':
    unittest.main()
