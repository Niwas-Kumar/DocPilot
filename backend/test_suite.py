# -*- coding: utf-8 -*-
"""DocPilot bug-hunt suite: pipeline edge cases, drive logic, API behavior.

Run with the server up on :8930 (API section skips politely if it is not).
"""
import io
import os
import sys
import time
import traceback

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import fitz
import requests as real_requests
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from backend.pipeline.daily_ca import process
from backend import drive

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORK = os.path.join(ROOT, "work", "tests")
EDGE = os.path.join(WORK, "edge_inputs")
os.makedirs(EDGE, exist_ok=True)

SAMPLES = os.environ.get("DOCPILOT_SAMPLES") or os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop", "UPSC sample")
EN_SAMPLE = os.path.join(SAMPLES, "Daily Current Affairs Summary (English)", "17th August 2026.docx")
HI_SAMPLE = os.path.join(SAMPLES, "Daily Current Affairs Summary (Hindi)", "Input 17 August-Daily News Hindi.docx")

PASS, FAIL = [], []


def check(name, fn):
    try:
        fn()
        PASS.append(name)
        print(f"  PASS  {name}")
    except Exception as e:
        FAIL.append((name, e))
        print(f"  FAIL  {name}: {type(e).__name__}: {e}")
        traceback.print_exc(limit=3)


def pdf_ok(path, min_pages=2):
    assert os.path.exists(path), "output missing"
    d = fitz.open(path)
    n = len(d)
    assert n >= min_pages, f"only {n} pages"
    # the first BODY page carries the number 2, the watermark and the footer
    # links. Since 2026-08-31 a front page precedes it whenever the file has
    # topic headings, and the cover prints no number - so look for the page
    # whose number line is exactly "2" (a substring test passed on the cover
    # by accident: "2026").
    def has_two(pg):
        return any(ln.strip() == "2" for ln in pg.get_text().split("\n"))
    p0 = next((d[i] for i in range(min(n, 2)) if has_two(d[i])), None)
    assert p0 is not None, "page number 2 not stamped on the first body page"
    assert p0.get_images(), "no images on the first body page (watermark/logo missing)"
    links = [l.get("uri", "") for l in p0.get_links()]
    assert any("whatsapp" in u for u in links), "whatsapp footer link missing"
    assert any("t.me" in u for u in links), "telegram footer link missing"
    d.close()
    return n


# ------------------------------------------------------------ edge builders

def build_minimal():
    p = os.path.join(EDGE, "minimal.docx")
    doc = Document()
    doc.add_paragraph("Just a plain paragraph without any structure.")
    doc.add_paragraph("Second paragraph with some more text to render.")
    doc.save(p)
    return p


def build_empty():
    p = os.path.join(EDGE, "empty.docx")
    Document().save(p)
    return p


def build_word_native():
    """Word-native constructs: built-in styles, bullet + DECIMAL lists, tables."""
    p = os.path.join(EDGE, "word_native.docx")
    doc = Document()
    doc.add_paragraph("Native Topic Heading", style="Heading 1")
    sub = doc.add_paragraph()
    sub.add_run("Subject: GS Paper 2: Governance").bold = True
    ctx = doc.add_paragraph()
    ctx.add_run("Context: ").bold = True
    ctx.add_run("Something happened that matters for the exam.")
    head = doc.add_paragraph()
    head.add_run("Fully Bold Section Heading:").bold = True
    doc.add_paragraph("Bullet item one", style="List Bullet")
    doc.add_paragraph("Bullet nested", style="List Bullet 2")
    doc.add_paragraph("Numbered step one", style="List Number")
    doc.add_paragraph("Numbered step two", style="List Number")
    t1 = doc.add_table(rows=1, cols=1)
    t1.cell(0, 0).text = "About Something: (UPSC CSE Prelims 2020)"
    t1.cell(0, 0).add_paragraph("Body line inside the box")
    t3 = doc.add_table(rows=2, cols=3)  # odd shape: neither infobox nor news
    for r in range(2):
        for c in range(3):
            t3.cell(r, c).text = f"cell {r}{c}"
    doc.add_paragraph("News In Shorts: Test Day", style="Title")
    t2 = doc.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "Topic One"
    t2.cell(0, 1).text = "Context: details about topic one."
    t2.cell(1, 0).text = "Topic Two"
    t2.cell(1, 1).text = "Context: details about topic two."
    doc.save(p)
    return p


def build_multisection():
    p = os.path.join(EDGE, "multisection.docx")
    from docx.enum.section import WD_SECTION
    doc = Document()
    doc.add_paragraph("Section one content.", style=None)
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Section two content after a section break.")
    doc.save(p)
    return p


def build_fake_docx():
    p = os.path.join(EDGE, "renamed_text.docx")
    with open(p, "w", encoding="utf-8") as f:
        f.write("this is not a zip archive")
    return p


# ------------------------------------------------------------ 1) pipeline

print("== 1. Pipeline edge cases ==")

check("english sample (baseline)", lambda: pdf_ok(
    process(EN_SAMPLE, os.path.join(WORK, "en"), "English"), min_pages=10))

# mechanical robustness only: Devanagari content must not crash the engine.
# (The OFFICIAL Hindi template differs from English - category is gated off
# until its own pipeline is built from the sample pair.)
check("hindi raw file processes without crashing", lambda: pdf_ok(
    process(HI_SAMPLE, os.path.join(WORK, "hi"), "Hindi"), min_pages=5))

check("minimal plain doc", lambda: pdf_ok(
    process(build_minimal(), os.path.join(WORK, "minimal"), "English")))

check("empty doc", lambda: pdf_ok(
    process(build_empty(), os.path.join(WORK, "empty"), "English")))

check("word-native styles + decimal list", lambda: pdf_ok(
    process(build_word_native(), os.path.join(WORK, "native"), "English")))


def _decimal_list_preserved():
    d = fitz.open(os.path.join(WORK, "native", "word_native_CA Analysis_UPSC_English.pdf"))
    alltext = "".join(p.get_text() for p in d)
    assert "1." in alltext or "1)" in alltext, "decimal numbering lost (turned into bullets?)"
    assert "Numbered step one" in alltext
    d.close()
check("decimal list stays numbered", _decimal_list_preserved)


def _multisection():
    out = process(build_multisection(), os.path.join(WORK, "multisec"), "English")
    d = fitz.open(out)
    assert len(d) >= 3, "expected >= 2 content pages + promo"
    for i in range(len(d) - 1):
        assert abs(d[i].rect.width - 612) < 1, f"page {i} not Letter width"
        assert d[i].get_images(), f"page {i} lost template images"
    d.close()
check("multi-section doc keeps template on all pages", _multisection)


def _fake_docx():
    try:
        process(build_fake_docx(), os.path.join(WORK, "fake"), "English")
        raise AssertionError("should have raised")
    except Exception as e:
        assert "Package" in str(e) or "not a zip" in str(e).lower() or "BadZip" in type(e).__name__
check("renamed .txt raises clean error", _fake_docx)

# ------------------------------------------------------------ 2) drive logic

print("== 2. Drive link logic ==")


def _parse_units():
    cases = {
        "https://docs.google.com/document/d/1AbCdEfGhIjKlMnOp_123-456/edit?usp=sharing": ("gdoc", "1AbCdEfGhIjKlMnOp_123-456"),
        "https://drive.google.com/file/d/1XyZ_abc-DEF456ghi/view?usp=drivesdk": ("file", "1XyZ_abc-DEF456ghi"),
        "https://drive.google.com/open?id=1OpenIdFile12345": ("file", "1OpenIdFile12345"),
        "https://drive.google.com/uc?export=download&id=1UcIdFile12345": ("file", "1UcIdFile12345"),
    }
    for url, expect in cases.items():
        assert drive.parse_drive_link(url) == expect, url
    for bad in ("https://example.com/x", "not a url", "https://drive.google.com/drive/folders/1abcdef1234567"):
        try:
            drive.parse_drive_link(bad)
            # folders parse as id= None? ensure garbage rejected
            if "folders" not in bad:
                raise AssertionError(f"accepted garbage: {bad}")
        except drive.DriveError:
            pass
check("parse_drive_link url shapes", _parse_units)


class FakeResp:
    def __init__(self, status=200, ctype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                 body=b"", cd=None, text=""):
        self.status_code = status
        self.headers = {"Content-Type": ctype}
        if cd:
            self.headers["Content-Disposition"] = cd
        self._body = body
        self.text = text

    def iter_content(self, n):
        for i in range(0, len(self._body), n):
            yield self._body[i:i + n]


def _mock_gdoc_download():
    real = drive.requests.Session
    docx_bytes = open(build_minimal(), "rb").read()

    class S:
        def get(self, url, **kw):
            assert "export?format=docx" in url
            return FakeResp(body=docx_bytes, cd="attachment; filename=\"My Doc.docx\"")
    drive.requests.Session = S
    try:
        p = drive.download_docx("https://docs.google.com/document/d/1AbCdEfGhIjKlMnOp_123-456/edit",
                                os.path.join(WORK, "dl1"))
        assert os.path.basename(p) == "My Doc.docx"
        assert os.path.getsize(p) == len(docx_bytes)
    finally:
        drive.requests.Session = real
check("mock: gdoc export happy path", _mock_gdoc_download)


def _mock_signin_page():
    real = drive.requests.Session

    class S:
        def get(self, url, **kw):
            return FakeResp(ctype="text/html; charset=utf-8", text="<html>Sign in</html>")
    drive.requests.Session = S
    try:
        try:
            drive.download_docx("https://docs.google.com/document/d/1AbCdEfGhIjKlMnOp_123-456/edit",
                                os.path.join(WORK, "dl2"))
            raise AssertionError("should raise DriveError")
        except drive.DriveError as e:
            assert "restricted" in str(e) or "sign-in" in str(e)
    finally:
        drive.requests.Session = real
check("mock: restricted link -> clear error", _mock_signin_page)


def _mock_interstitial():
    real = drive.requests.Session
    docx_bytes = open(build_minimal(), "rb").read()

    class S:
        def __init__(self):
            self.calls = 0

        def get(self, url, **kw):
            self.calls += 1
            if self.calls == 1:
                return FakeResp(ctype="text/html", text="...confirm=AbCd123...")
            assert "confirm=AbCd123" in url
            return FakeResp(body=docx_bytes, cd="attachment; filename=\"big.docx\"")
    drive.requests.Session = S
    try:
        p = drive.download_docx("https://drive.google.com/file/d/1XyZ_abc-DEF456ghi/view",
                                os.path.join(WORK, "dl3"))
        assert os.path.getsize(p) == len(docx_bytes)
    finally:
        drive.requests.Session = real
check("mock: big-file confirm interstitial", _mock_interstitial)


def _mock_cancel():
    real = drive.requests.Session
    docx_bytes = b"x" * (1024 * 1024)

    class S:
        def get(self, url, **kw):
            return FakeResp(body=docx_bytes)
    drive.requests.Session = S
    try:
        try:
            drive.download_docx("https://docs.google.com/document/d/1AbCdEfGhIjKlMnOp_123-456/e",
                                os.path.join(WORK, "dl4"), cancelled=lambda: True)
            raise AssertionError("should raise InterruptedError")
        except InterruptedError:
            pass
        leftovers = os.listdir(os.path.join(WORK, "dl4"))
        assert not leftovers, f"partial file left behind: {leftovers}"
    finally:
        drive.requests.Session = real
check("mock: cancel mid-download cleans up", _mock_cancel)

# ------------------------------------------------------------ 4) front pages
# No Word needed: extraction and date rules of the two cover stages.

print()
print("== 4. Front pages (extraction rules, no render) ==")


def _cover_en_extract():
    from backend.pipeline import cover
    if not os.path.exists(EN_SAMPLE):
        print("     (English sample not present - skipped)")
        return
    date, index, shorts = cover.extract_from_raw(EN_SAMPLE)
    assert date == "17 August 2026", date
    assert len(index) == 4 and index[0] == "Forest fires in India", index
    assert len(shorts) == 6 and shorts[0] == "SURYAPATH TIRANGA", shorts
    # typed block wins, empty block is extracted - independently
    d, i, sh, src = cover.resolve(EN_SAMPLE, date_in="", index_in="A\nB\n", shorts_in=None)
    assert (d, i, len(sh)) == ("17 August 2026", ["A", "B"], 6)
    assert src == {"date": "extracted", "index": "typed", "shorts": "extracted"}
check("cover EN: date/index/shorts extracted from the raw file", _cover_en_extract)


def _cover_en_dates():
    from backend.pipeline.cover import date_runs
    assert date_runs("17 August 2026") == [("17", False), ("th", True), (" August 2026", False)]
    assert date_runs("21 August 2026")[1] == ("st", True)
    assert date_runs("12 August 2026")[1] == ("th", True)
    assert date_runs("3rd Sept 2026")[1] == ("rd", True)
    assert date_runs("Special Edition") == [("Special Edition", False)]
check("cover EN: ordinal suffix rules", _cover_en_dates)


def _cover_hi_extract():
    from backend.pipeline import cover_hindi as ch
    if not os.path.exists(HI_SAMPLE):
        print("     (Hindi sample not present - skipped)")
        return
    index, bar, shorts = ch.extract_from_raw(HI_SAMPLE)
    assert len(index) == 4 and index[0].startswith("राष्ट्रीय"), index
    assert bar == "संक्षेप में समाचार", bar
    assert len(shorts) == 5 and shorts[-1] == "भारत-केन्या संबंध", shorts
    # the in-table Heading 2 cells (the shorts) must NOT leak into the index
    assert not any(s in index for s in shorts)
check("cover HI: topics = Heading 2 outside tables, shorts = news table", _cover_hi_extract)


def _cover_hi_dates():
    from backend.pipeline.cover_hindi import hindi_date
    for typed in ("17 August 2026", "17th Aug 2026", "17 अगस्त 2026", "17/08/2026",
                  "2026-08-17", "१७ अगस्त २०२६"):
        assert hindi_date(typed) == ("17 अगस्त 2026", True), typed
    assert hindi_date("3 Sept 2026") == ("3 सितंबर 2026", True)
    assert hindi_date("4 जुलाई 2026") == ("4 जुलाई 2026", True)
    assert hindi_date("Special") == ("Special", False)        # verbatim, flagged
    assert hindi_date("") == ("", False)
check("cover HI: typed date -> Hindi month, never invented", _cover_hi_dates)


def _cover_hi_fonts():
    from backend.pipeline import cover_hindi as ch
    fonts, note = ch.resolve_fonts()
    assert set(fonts) == {"title", "text", "heading"}
    assert fonts["title"] in ch.TITLE_CHAIN and fonts["text"] in ch.TEXT_CHAIN
    assert note.startswith("Cover fonts:")
check("cover HI: font chains resolve to installed families", _cover_hi_fonts)


def _cover_hi_build():
    """Fill the template without Word: one page section, Letter, all boxes filled."""
    from backend.pipeline import cover_hindi as ch
    from docx import Document
    out = os.path.join(WORK, "cover_hi_build.docx")
    numbers = {"क": 2, "ख": 5, "संक्षेप में समाचार": 9, "ग": 10}
    ch.build_cover_docx(out, "17 अगस्त 2026", ["क", "ख"], "संक्षेप में समाचार", ["ग", "घ"],
                        numbers, ch.INDEX_STEPS[0], {"title": "Nirmala UI", "text": "Nirmala UI",
                                                     "heading": "Nirmala UI"})
    doc = Document(out)
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    assert sect.find(qn("w:pgSz")).get(qn("w:h")) == "15840", "not Letter"
    paras = [p for p in body if p.tag == qn("w:p")]
    assert len(paras) == 1, f"{len(paras)} host paragraphs (must be 1)"
    texts = ["".join(t.text or "" for t in tb.iter(qn("w:t"))) for tb in body.iter(qn("w:txbxContent"))]
    assert any("17 अगस्त 2026" in t for t in texts), "date not filled"
    assert any("विषय सूची" in t and "घ" in t and "10" in t for t in texts), "index not filled"
    assert not any("जुलाई" in t for t in texts), "template sample date leaked"
    assert not list(body.iter("{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"))
check("cover HI: template fill (Letter, one host paragraph, no sample text)", _cover_hi_build)


def _xylem_furnish():
    """Xylem furniture swap, no Word: green border, 42.5pt top, strip footer
    with the PAGE field, numbering from 2, no English parts left behind."""
    from backend.pipeline import daily_ca_xylem as dx
    from backend.pipeline.styler import restyle_docx
    import zipfile, re
    src = EN_SAMPLE if os.path.exists(EN_SAMPLE) else build_minimal()
    styled = os.path.join(WORK, "xylem_furnish.docx")
    restyle_docx(src, styled, banner_text="Topics")
    dx.furnish(styled)
    doc = Document(styled)
    sect = doc.element.body.find(qn("w:sectPr"))
    assert sect.find(qn("w:pgMar")).get(qn("w:top")) == "850", "top margin"
    assert sect.find(qn("w:pgBorders")).find(qn("w:top")).get(qn("w:color")) == "00B050", "border colour"
    assert sect.find(qn("w:pgNumType")).get(qn("w:start")) == "2", "numbering start"
    assert sect.find(qn("w:titlePg")) is None
    refs = [(c.tag.split("}")[1], c.get(qn("w:type"))) for c in sect if c.tag.endswith("Reference")]
    assert sorted(refs) == [("footerReference", "default"), ("headerReference", "default")], refs
    z = zipfile.ZipFile(styled)
    footers = [n for n in z.namelist() if re.match(r"word/footer\d+\.xml", n)]
    assert len(footers) == 1, footers
    f = z.read(footers[0]).decode("utf-8")
    assert "PAGE" in f and "Impact" in f and "00B050" in f, "footer strip not carried"
    assert "whatsapp" not in f.lower(), "English link buttons leaked into the footer"
    big = [n for n in z.namelist() if n.startswith("word/media") and z.getinfo(n).file_size > 5_000_000]
    assert not big, f"leftover first-page watermark travelled: {big}"
check("xylem: furniture swapped from the template's section 2", _xylem_furnish)


def _xylem_cover_build():
    from backend.pipeline import cover_xylem as cx
    out = os.path.join(WORK, "xylem_cover_build.docx")
    filled = cx.build_cover_docx(out, "17 August 2026", ["A", "B"], ["S1"], cx.INDEX_SIZE_STEPS[0])
    assert filled == {"date": 2, "index": 2}, filled
    doc = Document(out)
    body = doc.element.body
    texts = ["".join(t.text or "" for t in tb.iter(qn("w:t"))) for tb in body.iter(qn("w:txbxContent"))]
    assert any("17" in t and "August 2026" in t for t in texts), "date not filled"
    assert any("News in Shorts" in t and "S1" in t for t in texts), "index not filled"
    assert not any("28th" in t or "Natural Outbursts" in t for t in texts), "template sample text leaked"
    assert len(body.findall(".//" + qn("w:sectPr"))) == 2, "template sections must stay"
    if os.path.exists(EN_SAMPLE):
        d, i, sh, _ = cx.resolve(EN_SAMPLE)
        assert d == "17 August 2026" and len(i) == 4 and len(sh) == 6
check("xylem: cover template fill + extraction shared with English", _xylem_cover_build)


def _xylem_box_collisions():
    """The side-box collision detector: flags the DMF box that body text ran
    through in the first Xylem output, and nothing on clean booklets."""
    from backend.pipeline import daily_ca_xylem as dx
    glitch = os.path.join(os.path.expanduser("~"), "Downloads",
                          "17th August 2026 (1)_CA Analysis_Xylem_UPSC_English.pdf")
    clean = [EN_SAMPLE.replace("17th August 2026.docx",
                               "17th August 2026_CA Analysis_UPSC_English_ With Cover.pdf")]
    checked = 0
    if os.path.exists(glitch):
        hits = dx.box_collisions(glitch)
        assert hits and hits[0][1].startswith("aboutdistrict"), hits
        checked += 1
    for c in clean:
        if os.path.exists(c):
            assert dx.box_collisions(c) == [], c
            checked += 1
    if not checked:
        print("     (no PDFs on this machine to check - skipped)")
check("xylem: side-box collision detector", _xylem_box_collisions)


def _bundled_fonts_present():
    """The OFL fonts the Hindi design needs must ship with the app - the
    client's PC had neither, so the cover title fell back and overprinted the
    PW logo."""
    from backend import fonts as F
    for filename, family, value_name in F.BUNDLED:
        path = os.path.join(F.FONT_DIR, filename)
        assert os.path.exists(path), f"bundled font missing: {filename}"
        assert os.path.getsize(path) > 50_000, filename
        assert value_name.endswith("(TrueType)"), value_name
    fams = {fam for _, fam, _ in F.BUNDLED}
    # Karma + Bakbak One: the Hindi Daily Editorial covers (2026-09-21)
    assert fams == {"Baloo 2", "Rozha One", "Karma", "Bakbak One"}, fams
    # the licence must travel with every family's fonts
    lic = {f.upper() for f in os.listdir(F.FONT_DIR) if f.upper().startswith("OFL")}
    for want in ("OFL-BALOO2.TXT", "OFL-ROZHAONE.TXT", "OFL-KARMA.TXT", "OFL-BAKBAKONE.TXT"):
        assert want in lic, (want, lic)
check("fonts: Baloo 2, Rozha One, Karma, Bakbak One bundled with their licences", _bundled_fonts_present)


def _font_install_is_idempotent():
    """An already-installed family is left completely alone (machine-wide or
    per user), and repeat calls never change anything."""
    from backend import fonts as F
    # decision logic, against synthetic registries
    assert F._available("Baloo 2", {"baloo 2", "arial"})
    assert F._available("Baloo 2", {"baloo 2 bold"})          # "Family Style"
    # a .ttc registers as "A & B & C"; installed_families() splits it first
    assert F._available("Nirmala UI", {"nirmala ui", "nirmala ui bold"})
    # near-misses must NOT count as installed
    assert not F._available("Rozha One", {"rozha two", "rozhaone"})
    assert not F._available("Baloo 2", {"baloo", "baloo bhai 2"})
    # on this machine both are installed, so nothing may be touched
    first = F.ensure_fonts(force=True)
    again = F.ensure_fonts(force=True)
    assert [r[:2] for r in first] == [r[:2] for r in again], (first, again)
    for family, status, _ in first:
        assert status in ("present", "installed", "skipped"), (family, status)
    assert "Bundled fonts:" in F.summary(first)
check("fonts: installer skips what is already there and repeats cleanly", _font_install_is_idempotent)


def _font_install_writes_when_missing():
    """The install path itself: copy + register + no crash on a second run.
    Uses a throwaway registry value so the real font setup is never touched."""
    from backend import fonts as F
    import tempfile, winreg
    value_name = "DocPilot Test Face (TrueType)"
    tmp = tempfile.mkdtemp(prefix="docpilot_fonts_")
    try:
        dest = F._install_one(F.BUNDLED[0][0], value_name, tmp)
        assert os.path.exists(dest), dest
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, F.REG_KEY) as k:
            got, _ = winreg.QueryValueEx(k, value_name)
        assert got == dest, (got, dest)
        F._install_one(F.BUNDLED[0][0], value_name, tmp)      # again: no error
    finally:
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, F.REG_KEY, 0,
                                winreg.KEY_SET_VALUE) as k:
                winreg.DeleteValue(k, value_name)
        except OSError:
            pass
        import shutil as _sh
        _sh.rmtree(tmp, ignore_errors=True)
check("fonts: install path copies + registers, and re-runs cleanly", _font_install_writes_when_missing)


def _hindi_drops_google_tab_frontmatter():
    """A Google Docs *tabs* export puts "Tab 1" (often with a hard page break)
    at the top; it left a near-empty page before the first topic."""
    from backend.pipeline.daily_ca_hindi import _drop_tab_artifacts
    from backend.pipeline.ooxml import GOOGLE_TAB_RE as TAB_ARTIFACT_RE
    from docx import Document
    from docx.oxml import parse_xml
    NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    assert TAB_ARTIFACT_RE.match("Tab 1") and TAB_ARTIFACT_RE.match("tab  12")
    assert not TAB_ARTIFACT_RE.match("Tab of contents")
    assert not TAB_ARTIFACT_RE.match("Table 1")

    def build(*xmls):
        d = Document()
        b = d.element.body
        for x in xmls:
            b.insert(0, parse_xml(x)) if False else b.append(parse_xml(x))
        return d

    tab_pb = f'<w:p {NS}><w:r><w:t>Tab 1</w:t></w:r><w:r><w:br w:type="page"/></w:r></w:p>'
    head = f'<w:p {NS}><w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>विषय</w:t></w:r></w:p>'
    d = build(tab_pb, head)
    removed, breaks = _drop_tab_artifacts(d.element.body, d)
    assert removed == 1, removed
    assert not list(d.element.body.iter(qn("w:br"))), "page break survived"

    # a "Tab 2" INSIDE the article must survive - only front matter is touched
    d2 = build(head, f'<w:p {NS}><w:r><w:t>Tab 2</w:t></w:r></w:p>')
    removed2, _ = _drop_tab_artifacts(d2.element.body, d2)
    assert removed2 == 0, "swallowed content after the first heading"
    # a clean document is untouched
    d3 = build(head)
    assert _drop_tab_artifacts(d3.element.body, d3) == (0, 0)
check("hindi: Google Docs 'Tab N' front matter dropped, real content kept",
      _hindi_drops_google_tab_frontmatter)


def _english_drops_google_tab_frontmatter():
    """English and Xylem read the SAME raw file, so the tabs artifact reaches
    them too - the scan is shared with the Hindi styler so they cannot drift."""
    from backend.pipeline.ooxml import drop_google_tab_frontmatter
    from backend.pipeline import styler
    from docx import Document
    from docx.oxml import parse_xml
    NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    d = Document()
    b = d.element.body
    b.append(parse_xml(f'<w:p {NS}><w:r><w:t>Tab 1</w:t></w:r>'
                       f'<w:r><w:br w:type="page"/></w:r></w:p>'))
    b.append(parse_xml(f'<w:p {NS}><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
                       f'<w:r><w:t>Forest fires in India</w:t></w:r></w:p>'))
    removed, _ = drop_google_tab_frontmatter(b, lambda el: styler._style_name(el, d))
    assert removed == 1, removed
    assert not list(b.iter(qn("w:br"))), "the hard page break survived"
    texts = ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in b.iter(qn("w:p"))]
    assert "Forest fires in India" in texts, texts
    # the styler really calls it
    import inspect
    assert "drop_google_tab_frontmatter" in inspect.getsource(styler.restyle_docx)
check("english/xylem: same 'Tab N' front matter dropped (shared scan)",
      _english_drops_google_tab_frontmatter)


def _hindi_cover_title_clears_logo():
    """Both title fonts must clear the PW logo. The fallback numbers are
    measured; the old ones put the title 20pt above its own box."""
    from backend.pipeline import cover_hindi as ch
    from docx import Document
    from docx.oxml.ns import qn as _qn
    for font, expect_exact in (("Rozha One", True), ("Nirmala UI", False)):
        out = os.path.join(WORK, f"hi_title_{font.replace(' ', '')}.docx")
        ch.build_cover_docx(out, "4 सितंबर 2026", ["क"], "संक्षेप में समाचार", ["ग"],
                            {}, ch.INDEX_STEPS[0],
                            {"title": font, "text": "Nirmala UI", "heading": "Baloo 2"})
        doc = Document(out)
        boxes = ch._textboxes(doc.element.body) if hasattr(ch, "_textboxes") else []
        xml = "".join(
            __import__("lxml.etree", fromlist=["etree"]).tostring(tb, encoding="unicode")
            for tb in doc.element.body.iter(_qn("w:txbxContent"))
            if ch.TITLE_TEXT in "".join(t.text or "" for t in tb.iter(_qn("w:t"))))
        assert ch.TITLE_TEXT in xml, font
        # the fallback must use the taller measured line box, never the old 1062
        assert 'w:line="1062"' not in xml, "the overlapping fallback line box is back"
        assert ('w:line="1500"' if expect_exact else 'w:line="1600"') in xml, font
check("hindi cover: title line box measured for both fonts (no logo overlap)",
      _hindi_cover_title_clears_logo)


def _hindi_body_font_matches_the_sample():
    """The published booklet sets its body in Adobe Devanagari and its
    headings in Baloo 2 - the design chains must say exactly that, and the
    matcher must find the body face however the provider names it."""
    from backend.pipeline import daily_ca_hindi as dh
    assert dh.BODY_CHAIN[0] == "Adobe Devanagari", dh.BODY_CHAIN
    assert dh.DISPLAY_CHAIN[0] == "Baloo 2", dh.DISPLAY_CHAIN
    # Baloo 2 sits second: the client asked for it by name and the app ships
    # it, so a machine with no Devanagari licence still gets a good booklet.
    assert dh.BODY_CHAIN[1] == "Baloo 2", dh.BODY_CHAIN
    for avail, want, want_designed in (
            ({"adobe devanagari", "baloo 2", "nirmala ui"}, "Adobe Devanagari", True),
            ({"adobe devanagari regular", "nirmala ui"}, "Adobe Devanagari", True),
            ({"adobedevanagari", "nirmala ui"}, "adobedevanagari", True),   # Adobe CC
            ({"baloo 2", "nirmala ui"}, "Baloo 2", False),   # what we ship
            ({"nirmala ui"}, "Nirmala UI", False)):
        got, designed = dh._pick(dh.BODY_CHAIN, avail)
        assert (got, designed) == (want, want_designed), (avail, got, designed)
    # when the design face is missing the note must tell the user how to fix it
    dh._font_cache.clear()
    try:
        dh._font_cache["fonts"] = ("Baloo 2", "Nirmala UI", True, False)
        notes = []
        dh.resolve_fonts(notes)
        joined = " ".join(notes)
        assert "Adobe Devanagari" in joined and "Creative Cloud" in joined, joined
        assert "ships with DocPilot" in joined, joined
        dh._font_cache.clear()
        dh._font_cache["fonts"] = ("Baloo 2", "Adobe Devanagari", True, True)
        ok_notes = []
        dh.resolve_fonts(ok_notes)
        assert not any("Creative Cloud" in n for n in ok_notes), ok_notes
    finally:
        dh._font_cache.clear()
check("hindi: body = Adobe Devanagari / headings = Baloo 2, as in the sample",
      _hindi_body_font_matches_the_sample)


def _topic_page_breaks():
    """Client rule (2026-09-08): every topic opens a new page, so an edit in
    the Word file reflows only its own topic instead of shunting the rest of
    the booklet. Two boundaries back to back (the English 'Topics' banner and
    the first topic title) must NOT be split - that would strand the banner."""
    from backend.pipeline import fitter
    from backend.pipeline.fitter import (ENGLISH_SPEC, HINDI_SPEC, _blocks,
                                         _break_between_topics)
    from backend.pipeline.daily_ca_xylem import XYLEM_SPEC
    from docx import Document
    from docx.oxml import parse_xml
    NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    # one switch for every edition (2026-10-12: off, topics flow as in the
    # published samples; True restores the 8 Sept layout)
    for spec in (ENGLISH_SPEC, HINDI_SPEC, XYLEM_SPEC):
        assert spec.topic_page_break is fitter.TOPIC_PAGE_BREAK, spec.name

    def topic(text):      # an English boundary: red 16pt H1
        return parse_xml(
            f'<w:p {NS}><w:r><w:rPr><w:sz w:val="32"/>'
            f'<w:color w:val="C00000"/></w:rPr><w:t>{text}</w:t></w:r></w:p>')

    def body(text):
        return parse_xml(f'<w:p {NS}><w:r><w:rPr><w:sz w:val="22"/></w:rPr>'
                         f'<w:t>{text}</w:t></w:r></w:p>')

    d = Document()
    b = d.element.body
    for el in (topic("A"), body("a1"), body("a2"), topic("B"), body("b1"),
               topic("C"), body("c1")):
        b.append(el)
    added = _break_between_topics(b, ENGLISH_SPEC)
    assert added == 2, added                       # before B and C, not A
    blocks = _blocks(b)
    starts = [i for i, el in enumerate(blocks) if ENGLISH_SPEC.is_boundary_block(el)]
    first = blocks[starts[0]].find(qn("w:pPr"))
    assert first is None or first.find(qn("w:pageBreakBefore")) is None,         "the first topic must not start with a page break"
    for i in starts[1:]:
        ppr = blocks[i].find(qn("w:pPr"))
        assert ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None
    assert _break_between_topics(b, ENGLISH_SPEC) == 0, "must be idempotent"

    # adjacent boundaries (banner immediately followed by a topic) stay together
    d2 = Document(); b2 = d2.element.body
    for el in (topic("Topics"), topic("First"), body("x"), topic("Second"), body("y")):
        b2.append(el)
    assert _break_between_topics(b2, ENGLISH_SPEC) == 1, "banner was split off"
check("daily CA: every topic opens a new page (adjacent banners stay put)",
      _topic_page_breaks)


def _word_convert_retries():
    """A transient Word COM failure must not throw away a whole job - the
    Prelims booklet renders dozens of times and used to die on one hiccup."""
    from backend.pipeline import convert
    real = convert._export_once
    killer = convert._kill_orphan_word
    calls = {"n": 0}

    def flaky(dx, pdf):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("simulated COM hiccup")
        with open(pdf, "wb") as f:
            f.write(b"%PDF-1.4 fake")

    convert._export_once = flaky
    convert._kill_orphan_word = lambda: None
    try:
        out = os.path.join(WORK, "retry_probe.pdf")
        src = os.path.join(WORK, "retry_probe.docx")
        Document().save(src)
        convert.docx_to_pdf(src, out)
        assert calls["n"] == 2, calls
        # a real failure still raises, and names the file
        convert._export_once = lambda dx, pdf: None
        try:
            convert.docx_to_pdf(src, os.path.join(WORK, "never.pdf"))
            raise AssertionError("should have raised")
        except RuntimeError as e:
            assert "retry_probe" in str(e) and "attempts" in str(e), str(e)
    finally:
        convert._export_once = real
        convert._kill_orphan_word = killer
check("word: docx->pdf retries a transient COM failure", _word_convert_retries)


def _qb_oversize_message():
    """The Prelims booklet must name the question it cannot fit, not say
    'please report this file'."""
    from backend.pipeline.prelims_qb import _oversize_message
    msg = _oversize_message({57: {"en"}})
    assert "Q57" in msg and "English" in msg and "numbering" in msg, msg
    assert "report this file" not in msg
    bug = _oversize_message({})
    assert "bug" in bug.lower(), bug
check("prelims QB: oversize error names the question", _qb_oversize_message)


def _every_category_resolves():
    """Every registered category must map to an importable pipeline with a
    process() - the queue resolves names through an explicit whitelist in
    jobs.py, and a new category that is not listed there fails only at
    submit time ('Unknown pipeline'), which is how the Xylem one shipped."""
    from backend.categories import CATEGORIES
    from backend import jobs
    for cat in CATEGORIES:
        if not cat.get("ready"):
            continue
        mod = jobs._pipeline(cat["pipeline"])
        assert callable(getattr(mod, "process", None)), cat["id"]
check("every ready category resolves to a pipeline with process()", _every_category_resolves)


print()
print(f"local sections: {len(PASS)} passed, {len(FAIL)} failed")
if FAIL:
    print("FAILURES:")
    for n, e in FAIL:
        print(" -", n, "->", e)
