# -*- coding: utf-8 -*-
"""Build clean DOCX templates from the official sample outputs.

Keeps everything that defines the DESIGN (styles, numbering, headers, footers,
theme, fonts, section setup) and removes the sample's editorial content, so the
pipelines can pour new content into a genuine PW template.

Run once:  python tools\make_templates.py
"""
import os
import shutil
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEST = os.path.join(ROOT, "assets", "templates")
SAMPLES = os.environ.get("DOCPILOT_SAMPLES") or os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop", "UPSC sample")

# keep_sections = which of the original sectPr blocks to preserve, in order.
#   prelims_ab   sect0 = 1-column  (title + answer-key grid), header1/footer1
#                sect1 = 2-column  (solutions),               header2/footer2
#                sect2/3 exist only to re-balance the final columns - dropped.
#   qb_bilingual sect0 = 1-column title, sect1.. = 2-column snaking body.
#   mains_ab     single section with titlePg (page-1 border) - keep as is.
# The Question Booklet keeps its front matter (English cover + Test
# Instructions + title band) and back matter (the same pages in Hindi): they
# are template furniture, identical in every booklet, and only the test name /
# subject inside them changes - the pipeline rewrites those at build time.
# Everything between is carved out and replaced by QUESTION_SENTINEL.
QUESTION_SENTINEL = "[[QUESTIONS]]"

JOBS = [
    dict(name="prelims_ab.docx",
         src=os.path.join(SAMPLES, "Prelims Test Series (Answer Booklet English)",
                          "Output Test 08_A_PTS 2708_Indian Polity Part II_FLT_"
                          "Prelims Wallah Test Series_UPSC_English_Aniket (1).docx"),
         keep_sections=[0, 1]),
    dict(name="prelims_qb_bilingual.docx",
         src=os.path.join(SAMPLES, "Prelims Test Series (Question Booklet Bilingual)",
                          "Output Test 08_Q_PTS 2708_Indian Polity Part II_FLT_"
                          "Prelims Wallah Test Series_UPSC_Bilingual_Aniket.docx"),
         carve=dict(front_last="Test – 08", back_first="Prelims Test Series 1.0")),
    dict(name="mains_ab_hindi.docx",
         src=os.path.join(SAMPLES, "Mains Test Series (Answer Booklet Hindi)",
                          "Output SDM2719_ Indian society_Srijan (P+M)_Hindi_Ans Booklet (1).docx"),
         keep_sections=[0]),
]


def _inherit_refs(sect, source):
    """Copy header/footer references from `source` into `sect` if it has none.

    Word normally inherits these from the preceding section. Once the earlier
    sections are deleted there is nothing to inherit from, so every retained
    section has to carry its own references or the branding disappears.
    """
    have = {(c.tag, c.get(qn('w:type')))
            for c in sect
            if c.tag in (qn('w:headerReference'), qn('w:footerReference'))}
    insert_at = 0
    for c in source:
        if c.tag not in (qn('w:headerReference'), qn('w:footerReference')):
            continue
        key = (c.tag, c.get(qn('w:type')))
        if key in have:
            continue
        import copy
        sect.insert(insert_at, copy.deepcopy(c))
        insert_at += 1


def strip_body(path, out_path, keep_sections):
    """Remove all paragraphs/tables from the body, keep chosen section setups."""
    doc = Document(path)
    body = doc.element.body

    # every sectPr in document order: paragraph-level ones first, body-final last
    ordered = []
    for p in body.findall(qn('w:p')):
        ppr = p.find(qn('w:pPr'))
        if ppr is not None and ppr.find(qn('w:sectPr')) is not None:
            ordered.append(ppr.find(qn('w:sectPr')))
    final_sect = body.find(qn('w:sectPr'))
    if final_sect is not None:
        ordered.append(final_sect)

    wanted = [ordered[i] for i in keep_sections if i < len(ordered)]
    if not wanted:
        wanted = ordered[:1]

    # the first section carries the real header/footer references; make sure
    # every retained section repeats them
    for s in wanted[1:]:
        _inherit_refs(s, wanted[0])

    for child in list(body):
        if child.tag in (qn('w:p'), qn('w:tbl'), qn('w:sdt')):
            body.remove(child)
    if final_sect is not None and final_sect.getparent() is body:
        body.remove(final_sect)

    # all but the last become paragraph-level section breaks; the last becomes
    # the body-final sectPr
    for s in wanted[:-1]:
        p = body.makeelement(qn('w:p'), {})
        ppr = body.makeelement(qn('w:pPr'), {})
        ppr.append(s)
        p.append(ppr)
        body.append(p)
    body.append(body.makeelement(qn('w:p'), {}))
    body.append(wanted[-1])

    doc.save(out_path)


def carve_body(path, out_path, front_last, back_first):
    """Keep the sample's front and back matter, replace the content between
    them with a single sentinel paragraph.

    `front_last` is the text of the last block to keep at the front (the title
    band), `back_first` the text of the first block to keep at the back (the
    Hindi cover). A paragraph that carries a section break is emptied rather
    than deleted, so the column layout of the carved-out region survives.
    """
    doc = Document(path)
    body = doc.element.body
    blocks = [c for c in body if c.tag in (qn('w:p'), qn('w:tbl'))]

    def text(el):
        return "".join(t.text or "" for t in el.iter(qn('w:t'))).strip()

    front_end = back_start = None
    for i, el in enumerate(blocks):
        t = text(el)
        if front_end is None and t.startswith(front_last):
            front_end = i
        if t.startswith(back_first):
            back_start = i            # last occurrence wins
    if front_end is None or back_start is None or back_start <= front_end:
        raise ValueError(f"carve markers not found ({front_last!r} / {back_first!r})")

    # Advance to the section break that CLOSES the title-band section: the
    # sentinel has to sit after it, or the questions land in that 1-column
    # section instead of the 2-column question region that follows.
    for j in range(front_end + 1, back_start):
        el = blocks[j]
        ppr = el.find(qn('w:pPr')) if el.tag == qn('w:p') else None
        if ppr is not None and ppr.find(qn('w:sectPr')) is not None:
            front_end = j
            break

    sentinel = parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:r><w:t xml:space="preserve">{QUESTION_SENTINEL}</w:t></w:r></w:p>')
    blocks[front_end].addnext(sentinel)

    for el in blocks[front_end + 1:back_start]:
        ppr = el.find(qn('w:pPr')) if el.tag == qn('w:p') else None
        if ppr is not None and ppr.find(qn('w:sectPr')) is not None:
            for child in list(el):          # keep the break, drop the words
                if child.tag != qn('w:pPr'):
                    el.remove(child)
            for child in list(ppr):
                if child.tag != qn('w:sectPr'):
                    ppr.remove(child)
            continue
        body.remove(el)

    doc.save(out_path)
    return front_end, back_start


def purge_orphan_media(path):
    """Drop image parts (and their relationships) that the stripped body no
    longer uses. Header/footer images are kept: their own .rels are separate.
    """
    import re

    zin = zipfile.ZipFile(path)
    names = zin.namelist()
    parts = {n: zin.read(n) for n in names}
    zin.close()

    # relationship ids actually referenced by the (now empty) document body
    doc_xml = parts["word/document.xml"].decode("utf-8", "replace")
    used_ids = set(re.findall(r'r:(?:embed|link|id)="([^"]+)"', doc_xml))

    # rewrite word/_rels/document.xml.rels, dropping unused image relationships
    rels_name = "word/_rels/document.xml.rels"
    rels_xml = parts[rels_name].decode("utf-8", "replace")

    def keep_rel(m):
        tag = m.group(0)
        rid = re.search(r'Id="([^"]+)"', tag)
        typ = re.search(r'Type="([^"]+)"', tag)
        if not rid or not typ:
            return tag
        is_image = typ.group(1).endswith("/image")
        if is_image and rid.group(1) not in used_ids:
            return ""            # orphan -> remove the relationship
        return tag

    rels_xml = re.sub(r"<Relationship\b[^>]*/>", keep_rel, rels_xml)
    parts[rels_name] = rels_xml.encode("utf-8")

    # which media files are still targeted by ANY .rels file?
    targeted = set()
    for n, data in parts.items():
        if not n.endswith(".rels"):
            continue
        xml = data.decode("utf-8", "replace")
        for t in re.findall(r'Target="([^"]+)"', xml):
            base = os.path.basename(t)
            if base:
                targeted.add(base)

    keep, dropped = [], []
    for n in names:
        if n.startswith("word/media/") and os.path.basename(n) not in targeted:
            dropped.append(n)
        else:
            keep.append(n)

    # [Content_Types].xml can keep its defaults; unused extensions are harmless
    tmp = path + ".tmp"
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    for n in keep:
        zout.writestr(n, parts[n])
    zout.close()
    os.replace(tmp, path)
    return dropped


def main():
    os.makedirs(DEST, exist_ok=True)
    only = set(a for a in sys.argv[1:] if not a.startswith("-"))
    for job in JOBS:
        src, name = job["src"], job["name"]
        if only and name not in only and name.replace(".docx", "") not in only:
            continue
        if not os.path.exists(src):
            print(f"  MISSING SOURCE: {src}")
            continue
        out = os.path.join(DEST, name)
        if job.get("carve"):
            carve_body(src, out, **job["carve"])
        else:
            strip_body(src, out, job["keep_sections"])
        dropped = purge_orphan_media(out)
        d = Document(out)
        z = zipfile.ZipFile(out)
        hf = sorted(n for n in z.namelist()
                    if n.startswith("word/header") or n.startswith("word/footer"))
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        styles = len(d.styles)
        print(f"{name}")
        print(f"   from     : {os.path.basename(src)}")
        print(f"   size     : {os.path.getsize(src)//1024} KB -> {os.path.getsize(out)//1024} KB "
              f"({len(dropped)} orphan media dropped)")
        print(f"   sections : {len(d.sections)}   styles: {styles}")
        print(f"   headers  : {hf}")
        print(f"   media    : {len(media)} kept")
        print(f"   body     : {len(d.paragraphs)} paragraphs, {len(d.tables)} tables")
        print()


if __name__ == "__main__":
    main()
